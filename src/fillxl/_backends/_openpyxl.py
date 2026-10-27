"""Openpyxl backend implementation for fillxl."""

from __future__ import annotations

import logging
from typing import TYPE_CHECKING, Any

import openpyxl

from fillxl._adapters import TabularAdapter
from fillxl._cells import populate_named_cells
from fillxl._charts import (
    TableChangeSpecs,
    update_chart_data_ranges,
)
from fillxl._lossy import check_lossy
from fillxl._sheets import first_remaining, require_a_visible_sheet, warn_sheet_not_found
from fillxl._tables import populate_table
from fillxl._validation import ColumnValidationStrategy, validate_columns

if TYPE_CHECKING:
    from collections.abc import Sequence
    from pathlib import Path

    from fillxl._config import FillXLConfig
    from fillxl._lossy import OnLossy

logger = logging.getLogger(__name__)

_OPENPYXL_SUFFIXES = (".xlsx", ".xlsm")


class OpenpyxlBackend:
    """Backend implementation using openpyxl.

    This backend works cross-platform and does not require Microsoft Excel
    to be installed.  It uses the openpyxl library to read/write ``.xlsx``
    files directly.
    """

    def __init__(self) -> None:
        self._workbook: openpyxl.Workbook | None = None
        self._source_suffix: str | None = None

    def _require_workbook(self) -> openpyxl.Workbook:
        """Return the loaded workbook or raise.

        Raises:
            RuntimeError: If no workbook has been loaded yet.

        """
        if self._workbook is None:
            msg = "No workbook loaded — call load() first"
            raise RuntimeError(msg)
        return self._workbook

    # ── Backend protocol ────────────────────────────────────────────────────

    def load(self, path: Path) -> None:
        """Open an Excel workbook from *path*.

        Args:
            path: Path to the workbook file.

        """
        logger.debug("Loading workbook from %s", path)

        ext = path.suffix.lower()
        if ext not in _OPENPYXL_SUFFIXES:
            msg = (
                f"File type '{ext}' is not supported by the openpyxl backend "
                "(supported: .xlsx, .xlsm). Use engine='win32com' for .xlsb and .xls."
            )
            raise ValueError(msg)

        self._source_suffix = ext
        self._workbook = openpyxl.load_workbook(path, keep_vba=ext == ".xlsm", rich_text=True)
        _fix_chart_axes(self._workbook)

    def fill_named_cell(self, name: str, value: Any, config: FillXLConfig) -> None:  # noqa: ANN401
        """Write a scalar *value* to the named cell *name*.

        Args:
            name: Name of the named cell / defined name.
            value: Scalar value to write.
            config: Global fill configuration.

        """
        workbook = self._require_workbook()
        populate_named_cells(
            workbook,
            {name: value},
            escape_formulas=config.escape_formulas,
        )

    def fill_table(
        self,
        name: str,
        data: Any,  # noqa: ANN401
        config: FillXLConfig,
    ) -> TableChangeSpecs:
        """Fill the Excel table *name* with tabular *data*.

        Args:
            name: Name of the Excel table.
            data: Tabular data (DataFrame, list[dict], …).
            config: Global fill configuration.

        Returns:
            A :class:`TableChangeSpecs` describing the old and new table
            dimensions (so chart references can be updated).

        """
        workbook = self._require_workbook()

        # The adapter is built once and shared with populate_table so that
        # expensive data (e.g. a DuckDB relation) is only read once.
        adapter = TabularAdapter(data)
        _validate_table_columns(
            workbook=workbook,
            table_name=name,
            data_columns=adapter.columns,
            strategy=config.column_validation_strategy,
        )

        return populate_table(
            workbook,
            name,
            adapter,
            copy_style_in_tables=config.copy_table_styles,
            extend_conditional_formatting=config.extend_conditional_formatting,
            escape_formulas=config.escape_formulas,
        )

    def update_charts(self, spec: TableChangeSpecs) -> None:
        """Update chart data ranges after a table resize.

        Args:
            spec: The table dimension change to reflect in charts.

        """
        update_chart_data_ranges(spec, self._require_workbook())

    def set_pivot_refresh(self) -> None:
        """Mark all pivot caches for refresh on load."""
        for worksheet in self._require_workbook().worksheets:
            for pivot in worksheet._pivots:  # type: ignore[attr-defined, unused-ignore]
                pivot.cache.refreshOnLoad = True
        logger.debug("Set all pivot tables to refresh on load")

    def hide_sheets(self, prefix: str) -> None:
        """Hide visible sheets whose names start with *prefix*.

        Args:
            prefix: Sheet name prefix (e.g. ``"_"``).

        """
        workbook = self._require_workbook()
        visible = _visible_titles(workbook)
        to_hide = [title for title in visible if title.startswith(prefix)]
        if not to_hide:
            return
        require_a_visible_sheet(visible, to_hide, "Hiding")
        active = _active_title(workbook)
        for title in to_hide:
            workbook[title].sheet_state = "hidden"
            logger.debug("Hidden sheet '%s'", title)
        if active is None or active in to_hide:
            _activate(workbook, first_remaining(visible, to_hide))

    def delete_sheets(self, names: Sequence[str]) -> None:
        """Delete the sheets named in *names*.

        Args:
            names: Sheet names to remove.

        """
        workbook = self._require_workbook()
        present = [name for name in dict.fromkeys(names) if name in workbook.sheetnames]
        for name in names:
            if name not in workbook.sheetnames:
                warn_sheet_not_found(name, "nothing was deleted")
        if not present:
            return
        visible = _visible_titles(workbook)
        require_a_visible_sheet(visible, present, "Deleting")
        active = _active_title(workbook)
        for name in present:
            workbook.remove(workbook[name])
            logger.debug("Deleted sheet '%s'", name)
        # openpyxl keeps the active *index*, which now points at another sheet.
        if active is not None and active not in present:
            _activate(workbook, active)
        else:
            _activate(workbook, first_remaining(visible, present))

    def set_active_sheet(self, name: str) -> None:
        """Set the active (displayed) sheet.

        Args:
            name: Name of the sheet to activate.

        Raises:
            ValueError: If the sheet is hidden.

        """
        workbook = self._require_workbook()
        if name not in workbook.sheetnames:
            warn_sheet_not_found(name, "the workbook keeps its active sheet")
            return
        if workbook[name].sheet_state != "visible":
            msg = f"Sheet '{name}' is hidden and cannot be the active sheet."
            raise ValueError(msg)
        _activate(workbook, name)
        logger.debug("Set active sheet to '%s'", name)

    def save(self, path: Path) -> None:
        """Save the workbook to *path*.

        Args:
            path: Destination path.

        """
        workbook = self._require_workbook()
        suffix = path.suffix.lower()
        if suffix != self._source_suffix:
            msg = (
                f"The openpyxl backend cannot turn a '{self._source_suffix}' template "
                f"into a '{suffix}' file. Save to a '{self._source_suffix}' path, or use "
                "engine='win32com' to convert between formats."
            )
            raise ValueError(msg)
        _fix_chart_axes(workbook)
        # openpyxl writes no cached formula results, so ask Excel to recalculate on open.
        from openpyxl.workbook.properties import CalcProperties

        calculation = workbook.calculation
        if calculation is None:
            calculation = CalcProperties()
            workbook.calculation = calculation
        calculation.fullCalcOnLoad = True
        logger.debug("Saving workbook to %s", path)
        workbook.save(path)

    def close(self) -> None:
        """Close the workbook and release resources."""
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    def check_output(self, template: Path, output: Path, policy: OnLossy) -> None:
        """Report template features openpyxl did not write back."""
        check_lossy(template, output, policy)


def _active_title(workbook: openpyxl.Workbook) -> str | None:
    active = workbook.active
    return None if active is None else str(active.title)


def _visible_titles(workbook: openpyxl.Workbook) -> list[str]:
    return [name for name in workbook.sheetnames if workbook[name].sheet_state == "visible"]


def _activate(workbook: openpyxl.Workbook, title: str) -> None:
    """Make *title* the active sheet and the only selected tab.

    A second selected tab makes Excel open the workbook with sheets grouped.
    Indices are positions in ``sheetnames``, which includes chart sheets.
    """
    names = workbook.sheetnames
    index = names.index(title)
    workbook.active = index
    for position, name in enumerate(names):
        _set_tab_selected(workbook[name], selected=position == index)


def _set_tab_selected(sheet: Any, *, selected: bool) -> None:  # noqa: ANN401
    views = getattr(sheet, "views", None)  # Worksheet
    if views is None:
        views = getattr(sheet, "sheetViews", None)  # Chartsheet
    if views is None:
        return
    for view in views.sheetView:
        view.tabSelected = selected


def _validate_table_columns(
    workbook: openpyxl.Workbook,
    table_name: str,
    data_columns: list[str],
    strategy: ColumnValidationStrategy,
) -> None:
    """Validate columns of incoming data against the Excel table definition.

    Args:
        workbook: The openpyxl Workbook containing the table.
        table_name: Name of the table to validate against.
        data_columns: Column names from the incoming data.
        strategy: Validation strategy (see :func:`fillxl._validation.validate_columns`).

    Raises:
        ValueError: If the strategy requires exact match and columns don't align.

    """
    from fillxl._formulas import openpyxl_formula_columns
    from fillxl._tables import openpyxl_table_geometry
    from fillxl._utils import find_worksheet_and_table, table_column_name

    worksheet, table = find_worksheet_and_table(workbook, table_name)
    geometry = openpyxl_table_geometry(table)
    table_columns = [table_column_name(col) for col in table.tableColumns]
    formula_columns = list(
        openpyxl_formula_columns(worksheet, table, geometry.first_data_row, geometry.body_last_row)
    )
    validate_columns(
        table_columns=table_columns,
        data_columns=data_columns,
        strategy=strategy,
        formula_columns=formula_columns,
        table_name=table_name,
    )


def _fix_chart_axes(workbook: openpyxl.Workbook) -> None:
    """Repair chart axes that openpyxl loaded with default values.

    Axes are fixed by role (``x_axis``/``y_axis``), not by class, so scatter
    charts (two value axes) are handled.  Positions openpyxl never defaults to
    (``"r"``, ``"t"``) and secondary axes (``crosses="max"``) are left alone.
    """
    from openpyxl.chart import BarChart
    from openpyxl.chart.axis import NumericAxis, TextAxis
    from openpyxl.chart.data_source import NumFmt

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        for chart in getattr(ws, "_charts", []):
            horizontal = isinstance(chart, BarChart) and chart.barDir == "bar"
            roles = (
                (getattr(chart, "x_axis", None), "l" if horizontal else "b"),
                (getattr(chart, "y_axis", None), "b" if horizontal else "l"),
            )
            for ax, expected_pos in roles:
                if ax is None:
                    continue
                secondary = ax.crosses == "max"
                if ax.delete is None:
                    ax.delete = False
                if ax.tickLblPos is None:
                    ax.tickLblPos = "nextTo"
                if ax.crosses is None:
                    ax.crosses = "autoZero"
                if not secondary and ax.axPos in (None, "l", "b"):
                    ax.axPos = expected_pos
                if isinstance(ax, TextAxis):
                    if ax.numFmt is None:
                        ax.numFmt = NumFmt(formatCode="General", sourceLinked=True)
                    if ax.auto is None:
                        ax.auto = True
                    if ax.lblAlgn is None:
                        ax.lblAlgn = "ctr"
                    if ax.noMultiLvlLbl is None:
                        ax.noMultiLvlLbl = False
                elif isinstance(ax, NumericAxis) and ax.crossBetween is None:
                    ax.crossBetween = "midCat"
