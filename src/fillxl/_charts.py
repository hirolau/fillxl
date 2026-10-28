"""Chart data range update logic."""

from __future__ import annotations

import logging
import re
from typing import TYPE_CHECKING, Any, NamedTuple

if TYPE_CHECKING:
    from collections.abc import Generator

    from openpyxl.workbook.workbook import Workbook

from openpyxl.worksheet.cell_range import CellRange

logger = logging.getLogger(__name__)


class TableChangeSpecs(NamedTuple):
    """Specification for a table range change.

    All row and column numbers are absolute worksheet coordinates and
    ``header_row`` is the first row of the table reference, so the first data
    row is ``header_row + header_rows``.  ``old_last_row`` and
    ``new_last_row`` are the last rows of the table's *data body*, which is
    what chart data ranges are compared against.  A totals row, if the table
    has one, sits below ``new_last_row`` and is not included.  Empty data
    still leaves one blank body row (Excel requires it), so a table with no
    data rows has ``new_last_row == first_data_row``.

    Attributes:
        worksheet_title: Title of the worksheet containing the table.
        header_row: First row of the table reference (the header row when the
            table has one, otherwise the first data row).
        first_col: 1-based index of the table's first column.
        last_col: 1-based index of the table's last column.
        old_last_row: Last data row of the table before the fill.
        new_last_row: Last data row of the table after the fill.
        header_rows: 1 when the table has a header row, 0 when headers are
            hidden (Excel names those columns ``Column1``, ``Column2``, …).

    Examples:
        >>> spec = TableChangeSpecs("Sheet1", 12, 2, 13, 13, 20)
        >>> spec.first_data_row
        13
        >>> spec.old_ref
        'B12:M13'
        >>> spec.new_ref
        'B12:M20'
        >>> headerless = TableChangeSpecs("Sheet1", 12, 2, 13, 13, 20, header_rows=0)
        >>> headerless.first_data_row
        12
        >>> headerless.old_ref
        'B12:M13'

    """

    worksheet_title: str
    header_row: int
    first_col: int
    last_col: int
    old_last_row: int
    new_last_row: int
    header_rows: int = 1

    @property
    def first_data_row(self) -> int:
        """Row number of the first data row (below the header, if any)."""
        return self.header_row + self.header_rows

    @property
    def old_ref(self) -> str:
        """The table's reference before the fill (e.g. ``B12:M13``)."""
        return _format_ref(self.first_col, self.header_row, self.last_col, self.old_last_row)

    @property
    def new_ref(self) -> str:
        """The table's reference after the fill (e.g. ``B12:M20``)."""
        return _format_ref(self.first_col, self.header_row, self.last_col, self.new_last_row)


def _format_ref(first_col: int, first_row: int, last_col: int, last_row: int) -> str:
    """Format a 1-based rectangle as an A1-style reference."""
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"


def _unquote_sheet_name(sheet_name: str) -> str:
    """Strip the quoting Excel adds around sheet names.

    Examples:
        >>> _unquote_sheet_name("'My Sheet'")
        'My Sheet'
        >>> _unquote_sheet_name("Data")
        'Data'

    """
    if len(sheet_name) >= 2 and sheet_name.startswith("'") and sheet_name.endswith("'"):
        return sheet_name[1:-1].replace("''", "'")
    return sheet_name


# A reference such as Data!$B$13:$B$20, 'My Sheet'!$B$13 or Försäljning!B13.
# The sheet is quoted ('' escapes a quote) or a run of characters that cannot
# appear in an unquoted name.  The look-behind stops a match from starting in
# the middle of a token such as [1]Data (a sheet in another workbook); the
# look-ahead stops $B$13 from matching the start of $B$134.
_SHEET_PATTERN = r"(?:'(?:[^']|'')+'|[^\s'!(),=\[\]{}\"]+)"
_CELL_PATTERN = r"\$?[A-Z]{1,3}\$?\d+"
REFERENCE_RE = re.compile(
    rf"(?<![^\s(),=])(?P<sheet>{_SHEET_PATTERN})!"
    rf"(?P<area>{_CELL_PATTERN}(?::{_CELL_PATTERN})?)(?![\w$:])"
)
_CELL_PARTS_RE = re.compile(r"(\$?[A-Z]{1,3}\$?)(\d+)")


def _with_last_row(area: str, last_row: int) -> str:
    """*area* as a range that ends on *last_row*, keeping its ``$`` markers.

    Examples:
        >>> _with_last_row("$B$13:$B$13", 20)
        '$B$13:$B$20'
        >>> _with_last_row("$B$13", 20)
        '$B$13:$B$20'
        >>> _with_last_row("B13:C13", 15)
        'B13:C15'

    """
    first, _, last = area.partition(":")
    end = _CELL_PARTS_RE.fullmatch(last or first)
    if end is None:  # pragma: no cover - REFERENCE_RE only matches valid cells
        msg = f"Not a cell reference: {area!r}"
        raise ValueError(msg)
    return f"{first}:{end.group(1)}{last_row}"


def update_reference_text(text: str, spec: TableChangeSpecs) -> str | None:
    """Rewrite references to the table's data body in *text*; ``None`` if none matched.

    A reference matches when it is on the table's sheet, starts on the first
    data row, ends on the old last data row and stays inside the table's
    columns.  Any number of references may appear, so non-contiguous series
    such as ``(Data!$B$13,Data!$D$13)`` work.

    Examples:
        >>> spec = TableChangeSpecs("Data", 12, 2, 13, 13, 20)
        >>> update_reference_text("Data!$B$13:$B$13", spec)
        'Data!$B$13:$B$20'
        >>> update_reference_text("Data!$B$13", spec)
        'Data!$B$13:$B$20'
        >>> update_reference_text("(Data!$B$13,Data!$D$13)", spec)
        '(Data!$B$13:$B$20,Data!$D$13:$D$20)'
        >>> update_reference_text("Data2!$B$13:$B$13", spec) is None
        True
        >>> update_reference_text("Data!$B$12:$B$13", spec) is None
        True
        >>> update_reference_text("[1]Data!$B$13", spec) is None
        True
        >>> swedish = TableChangeSpecs("Försäljning", 1, 1, 2, 2, 5)
        >>> update_reference_text("Försäljning!$B$2", swedish)
        'Försäljning!$B$2:$B$5'
        >>> apostrophe = TableChangeSpecs("Bob's Data", 1, 1, 2, 2, 5)
        >>> update_reference_text("'Bob''s Data'!$B$2", apostrophe)
        "'Bob''s Data'!$B$2:$B$5"

    """
    if spec.new_last_row == spec.old_last_row:
        return None
    changed = False

    def replace(match: re.Match[str]) -> str:
        nonlocal changed
        if _unquote_sheet_name(match["sheet"]) != spec.worksheet_title:
            return match[0]
        bounds = CellRange(match["area"])
        if bounds.min_row != spec.first_data_row or bounds.max_row != spec.old_last_row:
            return match[0]
        if bounds.min_col < spec.first_col or bounds.max_col > spec.last_col:
            return match[0]
        changed = True
        return f"{match['sheet']}!{_with_last_row(match['area'], spec.new_last_row)}"

    updated = REFERENCE_RE.sub(replace, text)
    return updated if changed else None


def point_count(text: str) -> int:
    """Number of cells the references in *text* cover.

    Examples:
        >>> point_count("(Data!$B$13:$B$20,Data!$D$13:$D$20)")
        16

    """
    total = 0
    for match in REFERENCE_RE.finditer(text):
        size = CellRange(match["area"]).size
        total += size["rows"] * size["columns"]
    return total


def point_count_rows(text: str) -> int:
    """Number of rows covered by a chart reference, for multi-level category caches."""
    total = 0
    for match in REFERENCE_RE.finditer(text):
        size = CellRange(match["area"]).size
        total += size["rows"]
    return total


def split_formula_arguments(text: str) -> list[str]:
    """Split *text* on commas that are not inside quotes or parentheses.

    Examples:
        >>> split_formula_arguments('"a,b",(Data!$B$1,Data!$C$1),1')
        ['"a,b"', '(Data!$B$1,Data!$C$1)', '1']
        >>> split_formula_arguments("'Bob''s, Data'!$A$1,2")
        ["'Bob''s, Data'!$A$1", '2']

    """
    parts: list[str] = []
    depth = 0
    start = 0
    quote: str | None = None
    for index, char in enumerate(text):
        if quote is not None:
            if char == quote:  # a doubled quote closes and reopens: net no change
                quote = None
            continue
        if char in "'\"":
            quote = char
        elif char == "(":
            depth += 1
        elif char == ")":
            depth -= 1
        elif char == "," and depth == 0:
            parts.append(text[start:index])
            start = index + 1
    parts.append(text[start:])
    return parts


def all_chart_series(workbook: Workbook) -> Generator[Any, None, None]:
    """Yield all chart series from all worksheets in the workbook.

    Args:
        workbook: openpyxl Workbook instance

    Yields:
        Chart series objects

    """
    for sheet in workbook.sheetnames:
        for chart in workbook[sheet]._charts:
            plots = getattr(chart, "_charts", None) or [chart]
            for plot in plots:
                yield from plot.ser


def all_chart_series_parts(workbook: Workbook) -> Generator[Any, None, None]:
    """Yield every chart data-source object (``cat``, ``val``, …) of every series.

    The labels cover bubble charts (``bubbleSize``) as well as the usual
    category and value sources, so a table resize reaches every reference a
    chart holds.

    Args:
        workbook: openpyxl Workbook instance

    Yields:
        Chart ``*Ref``/``*Ser`` data-source objects

    """
    labels = ("cat", "val", "xVal", "yVal", "bubbleSize")
    for series in all_chart_series(workbook=workbook):
        for label in labels:
            part = getattr(series, label, None)
            if part is not None:
                yield part
        yield from _error_bar_parts(series)


def _error_bar_parts(series: Any) -> Generator[Any, None, None]:  # noqa: ANN401
    """The data sources of a series' custom error bars, if it has any."""
    for direction in ("errBars", "errBarsX", "errBarsY"):
        error_bars = getattr(series, direction, None)
        if error_bars is None:
            continue
        for side in ("plus", "minus"):
            part = getattr(error_bars, side, None)
            if part is not None:
                yield part


def _reference_formulas(part: Any) -> Generator[Any, None, None]:  # noqa: ANN401
    """Yield the reference object inside one data source.

    A multi-level category source holds its range in a ``multiLvlStrRef``
    rather than in a ``strRef``, so both spellings are looked for.
    """
    for ref_type in ("strRef", "numRef", "multiLvlStrRef"):
        reference = getattr(part, ref_type, None)
        if reference is not None:
            yield reference


def all_chart_series_references(workbook: Workbook) -> Generator[Any, None, None]:
    """Yield all chart series data references (strRef/numRef) from all series.

    Args:
        workbook: openpyxl Workbook instance

    Yields:
        Reference objects (strRef or numRef) that contain chart data ranges

    """
    for part in all_chart_series_parts(workbook):
        yield from _reference_formulas(part)


def update_chart_data_ranges(spec: TableChangeSpecs, workbook: Workbook) -> None:
    """Point chart series that show the table's data body at its new rows.

    Args:
        spec: TableChangeSpecs describing the old and new table dimensions
        workbook: openpyxl Workbook instance

    """
    for reference in all_chart_series_references(workbook):
        ref = reference.f
        if not ref:
            continue
        new_ref = update_reference_text(ref, spec)
        if new_ref is None:
            continue
        reference.f = new_ref
        # openpyxl does not update cached point counts; stale counts make Excel
        # drop axis labels (especially on area charts).
        for cache_name in ("strCache", "numCache", "multiLvlStrCache"):
            cache = getattr(reference, cache_name, None)
            if cache is None:
                continue
            if cache_name == "multiLvlStrCache":
                count = point_count_rows(new_ref)
            else:
                count = point_count(new_ref)
            cache.ptCount = count
        logger.debug("Updated chart reference: %s -> %s", ref, new_ref)
