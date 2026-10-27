"""Row arithmetic for filling an Excel table, shared by both backends.

The backends used to work out the rows to check, clear and resize on their
own, and parity bugs hid in the differences.  :func:`plan_table_fill` is pure,
so the arithmetic is tested once, without Excel.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from attrs import define
from openpyxl.utils import get_column_letter

from ._charts import TableChangeSpecs

if TYPE_CHECKING:
    from collections.abc import Iterable

RowSpan = tuple[int, int]
"""Inclusive ``(first_row, last_row)`` range of worksheet rows."""


@define(frozen=True, kw_only=True)
class BlockedRegion:
    """A worksheet rectangle that a growing table must not write into.

    Attributes:
        description: How to name the region in an error message, e.g.
            ``"table 'Other'"`` or ``"the merged range A5:C6"``.
        first_row: First row of the region.
        last_row: Last row of the region.
        first_col: First column of the region.
        last_col: Last column of the region.

    """

    description: str
    first_row: int
    last_row: int
    first_col: int
    last_col: int

    def overlaps(self, rows: RowSpan, first_col: int, last_col: int) -> bool:
        """Whether the rectangle overlaps *rows* between *first_col* and *last_col*."""
        first_row, last_row = rows
        return (
            self.first_col <= last_col
            and self.last_col >= first_col
            and self.first_row <= last_row
            and self.last_row >= first_row
        )


def blocked_area(
    rows: RowSpan,
    first_col: int,
    last_col: int,
    regions: Iterable[BlockedRegion],
) -> BlockedRegion | None:
    """The first region a table growing into *rows* would overlap, or ``None``.

    Examples:
        >>> merged = BlockedRegion(
        ...     description="the merged range C15:D16",
        ...     first_row=15,
        ...     last_row=16,
        ...     first_col=3,
        ...     last_col=4,
        ... )
        >>> blocked_area((14, 20), 1, 13, [merged]) is merged
        True
        >>> blocked_area((14, 20), 5, 6, [merged]) is None
        True
        >>> blocked_area((20, 22), 1, 13, [merged]) is None
        True

    """
    return next(
        (region for region in regions if region.overlaps(rows, first_col, last_col)),
        None,
    )


@define(frozen=True, kw_only=True)
class TableGeometry:
    """Where a table sits before the fill (1-based worksheet coordinates).

    Attributes:
        top_row: First row of the table: its header row, if it has one.
        last_row: Last row of the table: its totals row, if it has one.
        first_col: First column of the table.
        last_col: Last column of the table.
        header_rows: 1 with a header row, 0 when headers are hidden.
        totals_rows: 1 with a totals row, otherwise 0.

    """

    top_row: int
    last_row: int
    first_col: int
    last_col: int
    header_rows: int
    totals_rows: int

    @property
    def first_data_row(self) -> int:
        """First row of the data body."""
        return self.top_row + self.header_rows

    @property
    def body_last_row(self) -> int:
        """Last row of the data body (the totals row excluded)."""
        return self.last_row - self.totals_rows


@define(frozen=True, kw_only=True)
class TablePlan:
    """What filling a table with ``data_rows`` rows does to the worksheet.

    Attributes:
        geometry: The table before the fill.
        data_rows: Number of rows being written (0 clears the table).
        new_body_last_row: Last data row after the fill.  An Excel table needs
            one body row, so empty data keeps a single cleared row.
        new_last_row: Last row of the table after the fill, totals included.
        overflow_rows: Rows the table grows into; they must be empty.
        unused_body_rows: Old body rows that receive no data and are cleared.
        vacated_rows: Rows inside the old table but outside the new one.

    """

    geometry: TableGeometry
    data_rows: int
    new_body_last_row: int
    new_last_row: int
    overflow_rows: RowSpan | None
    unused_body_rows: RowSpan | None
    vacated_rows: RowSpan | None

    @property
    def first_data_row(self) -> int:
        """First row of the data body."""
        return self.geometry.first_data_row

    @property
    def resized(self) -> bool:
        """Whether the table's last row moves."""
        return self.new_last_row != self.geometry.last_row

    @property
    def new_ref(self) -> str:
        """The table's reference after the fill, e.g. ``B12:M20``."""
        g = self.geometry
        return (
            f"{get_column_letter(g.first_col)}{g.top_row}:"
            f"{get_column_letter(g.last_col)}{self.new_last_row}"
        )

    @property
    def new_autofilter_ref(self) -> str:
        """The autofilter range after the fill (Excel excludes the totals row)."""
        g = self.geometry
        return (
            f"{get_column_letter(g.first_col)}{g.top_row}:"
            f"{get_column_letter(g.last_col)}{self.new_body_last_row}"
        )

    def change_spec(self, worksheet_title: str) -> TableChangeSpecs:
        """The change in chart terms, for updating chart references."""
        g = self.geometry
        return TableChangeSpecs(
            worksheet_title=worksheet_title,
            header_row=g.top_row,
            first_col=g.first_col,
            last_col=g.last_col,
            old_last_row=g.body_last_row,
            new_last_row=self.new_body_last_row,
            header_rows=g.header_rows,
        )


def plan_table_fill(geometry: TableGeometry, data_rows: int) -> TablePlan:
    """Plan filling the table at *geometry* with *data_rows* rows.

    Examples:
        A one-row table that receives eight rows grows into rows 14-20:

        >>> geometry = TableGeometry(
        ...     top_row=12, last_row=13, first_col=2, last_col=13, header_rows=1, totals_rows=0
        ... )
        >>> plan = plan_table_fill(geometry, 8)
        >>> plan.new_ref, plan.overflow_rows, plan.unused_body_rows, plan.vacated_rows
        ('B12:M20', (14, 20), None, None)

        A table with a totals row shrinks from seven body rows to three:

        >>> geometry = TableGeometry(
        ...     top_row=12, last_row=20, first_col=2, last_col=13, header_rows=1, totals_rows=1
        ... )
        >>> plan = plan_table_fill(geometry, 3)
        >>> plan.new_body_last_row, plan.new_last_row, plan.unused_body_rows, plan.vacated_rows
        (15, 16, (16, 19), (17, 20))

        Empty data keeps one cleared body row:

        >>> geometry = TableGeometry(
        ...     top_row=12, last_row=20, first_col=2, last_col=13, header_rows=1, totals_rows=0
        ... )
        >>> plan = plan_table_fill(geometry, 0)
        >>> plan.new_body_last_row, plan.unused_body_rows, plan.vacated_rows
        (13, (13, 20), (14, 20))

    """
    if data_rows < 0:
        msg = f"data_rows must not be negative, got {data_rows}"
        raise ValueError(msg)
    g = geometry
    new_body_last_row = g.first_data_row + max(data_rows, 1) - 1
    new_last_row = new_body_last_row + g.totals_rows
    first_unused = g.first_data_row + data_rows
    return TablePlan(
        geometry=g,
        data_rows=data_rows,
        new_body_last_row=new_body_last_row,
        new_last_row=new_last_row,
        overflow_rows=(g.last_row + 1, new_last_row) if new_last_row > g.last_row else None,
        unused_body_rows=(
            (first_unused, g.body_last_row) if first_unused <= g.body_last_row else None
        ),
        vacated_rows=(new_last_row + 1, g.last_row) if new_last_row < g.last_row else None,
    )
