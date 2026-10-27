"""Table population logic."""

from __future__ import annotations

import logging
from copy import copy
from typing import TYPE_CHECKING, Any, NamedTuple

from openpyxl.cell.cell import Cell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange

from ._adapters import TabularAdapter
from ._charts import TableChangeSpecs
from ._formulas import openpyxl_formula_columns, warn_about_ignored_formula_data
from ._normalize import (
    check_cell_value,
    check_table_cell_scalar,
    is_text_cell_value,
    prepare_cell_value,
)
from ._plan import (
    BlockedRegion,
    RowSpan,
    TableGeometry,
    blocked_area,
    plan_table_fill,
)
from ._utils import find_worksheet_and_table, table_column_name

if TYPE_CHECKING:
    from collections.abc import Sequence

    from openpyxl.workbook.workbook import Workbook
    from openpyxl.worksheet.table import Table
    from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


def is_single_cell(cell: object) -> bool:
    """Check if an object is a single Cell instance (not a MergedCell).

    Args:
        cell: Object to check

    Returns:
        True if *cell* is a :class:`~openpyxl.cell.cell.Cell` instance.

    Examples:
        >>> from openpyxl.cell.cell import Cell
        >>> is_single_cell(Cell(None, None, None, None))
        True
        >>> is_single_cell("not a cell")
        False
        >>> from openpyxl.worksheet.cell_range import CellRange
        >>> is_single_cell(CellRange("A1:B2"))
        False

    """
    return isinstance(cell, Cell)


def openpyxl_table_geometry(table: Table) -> TableGeometry:
    """Where *table* sits, read from its reference and header/totals settings."""
    ref = CellRange(table.ref)
    return TableGeometry(
        top_row=ref.min_row,
        last_row=ref.max_row,
        first_col=ref.min_col,
        last_col=ref.max_col,
        header_rows=0 if table.headerRowCount == 0 else 1,
        totals_rows=_totals_row_count(table),
    )


def _clear_rows(
    worksheet: Worksheet,
    rows: RowSpan,
    first_col: int,
    last_col: int,
    *,
    keep: Sequence[int] = (),
) -> None:
    """Clear values (not styles) in *rows* between *first_col* and *last_col*.

    Columns in *keep* are left alone: a calculated column keeps its formula,
    which is what Excel does when it shrinks a table.
    """
    first_row, last_row = rows
    skipped = set(keep)
    for row in range(first_row, last_row + 1):
        for col in range(first_col, last_col + 1):
            if col in skipped:
                continue
            cell = worksheet._cells.get((row, col))
            if cell is not None and is_single_cell(cell):
                cell.value = None


def _check_overflow_region(
    worksheet: Worksheet,
    table_name: str,
    table: Table,
    rows: RowSpan,
    first_col: int,
    last_col: int,
    data_rows: int,
) -> None:
    """Refuse to grow a table into another table or into merged cells.

    Those cells can look empty and still break the file: an overlapping table
    makes Excel repair the workbook, and writing to a merged cell raises
    openpyxl's "read-only" ``AttributeError``.
    """
    blocked = blocked_area(
        rows,
        first_col,
        last_col,
        _occupied_regions(worksheet, table),
    )
    if blocked is None:
        return
    msg = (
        f"Table '{table_name}' cannot expand to {data_rows} rows: "
        f"{blocked.description} is in the way. Move or remove the content below "
        "the table first."
    )
    raise ValueError(msg)


def _occupied_regions(worksheet: Worksheet, table: Table) -> list[BlockedRegion]:
    """Everything on *worksheet* that a table must not grow into."""
    regions = [
        BlockedRegion(
            description=f"table {other.name!r} at {other.ref}",
            first_row=CellRange(other.ref).min_row,
            last_row=CellRange(other.ref).max_row,
            first_col=CellRange(other.ref).min_col,
            last_col=CellRange(other.ref).max_col,
        )
        for other in worksheet.tables.values()
        if other is not table
    ]
    regions.extend(
        BlockedRegion(
            description=f"the merged range {merged}",
            first_row=merged.min_row,
            last_row=merged.max_row,
            first_col=merged.min_col,
            last_col=merged.max_col,
        )
        for merged in worksheet.merged_cells.ranges
    )
    return regions


def _check_overflow(
    worksheet: Worksheet,
    table_name: str,
    rows: RowSpan,
    first_col: int,
    last_col: int,
    data_rows: int,
) -> None:
    """Refuse to grow a table over non-empty cells.

    ``_cells`` is read directly: ``worksheet.cell()`` would create the cell.
    """
    first_row, last_row = rows
    for row in range(first_row, last_row + 1):
        for col in range(first_col, last_col + 1):
            cell = worksheet._cells.get((row, col))
            if cell is not None and cell.value is not None:
                msg = (
                    f"Table '{table_name}' cannot expand to {data_rows} rows: "
                    f"cell {get_column_letter(col)}{row} already contains data "
                    f"({cell.value!r}). Move or remove the content below the table first."
                )
                raise ValueError(msg)


def populate_table(
    workbook: Workbook,
    table_name: str,
    adapter: TabularAdapter,
    *,
    copy_style_in_tables: bool = True,
    extend_conditional_formatting: bool = True,
    escape_formulas: bool = False,
) -> TableChangeSpecs:
    """Populate an Excel table with tabular data.

    Args:
        workbook: openpyxl Workbook instance
        table_name: Name of the table to populate
        adapter: Adapter holding the column names and rows to write.  The
            caller builds it (and reuses it for column validation) so that
            the underlying data is not read twice.
        copy_style_in_tables: If True, copy cell styles from first data row to new rows
        extend_conditional_formatting: If True, extend conditional formatting
            rules to cover new rows
        escape_formulas: If True, strings starting with ``=`` are written as
            literal text instead of becoming formulas.

    Returns:
        TableChangeSpecs describing the old and new table dimensions

    Raises:
        KeyError: If table is not found
        ValueError: If the table cannot expand into non-empty cells.

    """
    columns = adapter.columns
    rows = adapter.rows

    worksheet, table = find_worksheet_and_table(workbook, table_name)
    geometry = openpyxl_table_geometry(table)
    plan = plan_table_fill(geometry, len(rows))
    first_col, last_col = geometry.first_col, geometry.last_col
    header_rows = geometry.header_rows

    # Calculated columns must be detected before any body cell is cleared.
    formula_columns = openpyxl_formula_columns(
        worksheet, table, geometry.first_data_row, geometry.body_last_row
    )
    warn_about_ignored_formula_data(table_name, columns, rows, formula_columns)

    # openpyxl writes cells in place, so the totals row is captured before
    # clearing and recreated at its new position afterwards.
    totals_row_cells = (
        _read_row(worksheet, geometry.last_row, first_col, last_col)
        if geometry.totals_rows
        else None
    )

    if plan.overflow_rows is not None:
        _check_overflow_region(
            worksheet,
            table_name,
            table,
            plan.overflow_rows,
            first_col,
            last_col,
            len(rows),
        )
        _check_overflow(worksheet, table_name, plan.overflow_rows, first_col, last_col, len(rows))

    # The autoFilter must stay in step with the table, or Excel repairs the file.
    original_ref = table.ref
    table.ref = plan.new_ref
    if table.autoFilter is not None:
        table.autoFilter.ref = plan.new_autofilter_ref
    logger.debug("Table '%s' resized: %s -> %s", table_name, original_ref, plan.new_ref)

    table_first_cell = worksheet.cell(row=geometry.top_row, column=first_col)
    if not is_single_cell(table_first_cell):
        msg = "Expected a single cell! This should never happen in a table."
        raise ValueError(msg)

    # Calculated columns keep their formulas, including when the data is empty.
    formula_column_offsets = {
        offset
        for offset, column in enumerate(table.tableColumns)
        if table_column_name(column) in formula_columns
    }
    keep_columns = [first_col + offset for offset in formula_column_offsets]
    if plan.unused_body_rows is not None:
        _clear_rows(worksheet, plan.unused_body_rows, first_col, last_col, keep=keep_columns)
    if geometry.totals_rows:
        _clear_rows(worksheet, (geometry.last_row, geometry.last_row), first_col, last_col)

    data_column_map = {str(name): idx for idx, name in enumerate(columns)}
    warned: set[tuple[str | None, str | None]] = set()

    for col_ix, column in enumerate(table.tableColumns):
        first_row_style = table_first_cell.offset(column=col_ix, row=header_rows)._style
        column_name = table_column_name(column)

        if column_name in formula_columns:
            formula = formula_columns[column_name]
            for row_ix in range(plan.data_rows):
                cell = table_first_cell.offset(column=col_ix, row=header_rows + row_ix)
                if copy_style_in_tables:
                    cell._style = copy(first_row_style)
                cell.value = formula

        elif column_name in data_column_map:
            data_col_idx = data_column_map[column_name]
            for row_ix, row_values in enumerate(rows):
                cell = table_first_cell.offset(column=col_ix, row=header_rows + row_ix)
                raw_value = row_values[data_col_idx]
                check_table_cell_scalar(
                    raw_value, table=table_name, row=row_ix + 1, column=column_name
                )
                # Style first: a date format must be in place before the value
                # arrives, or openpyxl replaces it with its default date format.
                if copy_style_in_tables:
                    cell._style = copy(first_row_style)
                value = prepare_cell_value(raw_value)
                check_cell_value(
                    value, table=table_name, row=row_ix + 1, column=column_name, warned=warned
                )
                cell.value = value
                # Force text for strings openpyxl would reinterpret: error
                # values such as "#N/A", and formulas when escaping is on.
                if is_text_cell_value(value, escape_formulas=escape_formulas):
                    cell.data_type = "s"

        else:
            # The data has no such column: clear the template's sample values.
            column_index = first_col + col_ix
            _clear_rows(
                worksheet,
                (plan.first_data_row, plan.new_body_last_row),
                column_index,
                column_index,
            )

    if totals_row_cells is not None:
        _write_row(worksheet, plan.new_last_row, first_col, totals_row_cells)
        if plan.vacated_rows is not None:
            # The old totals row keeps its bold text and rules unless the style
            # goes too; Excel moves the row up and leaves a plain row behind.
            _reset_row_style(
                worksheet,
                geometry.last_row,
                plan.first_data_row,
                first_col,
                last_col,
            )

    if extend_conditional_formatting:
        _extend_conditional_formatting(
            worksheet=worksheet,
            table_start_column=first_col,
            table_end_column=last_col,
            first_data_row=plan.first_data_row,
            table_last_row=plan.new_body_last_row,
        )

    return plan.change_spec(str(worksheet.title))


def _reset_row_style(
    worksheet: Worksheet,
    row: int,
    template_row: int,
    first_col: int,
    last_col: int,
) -> None:
    """Give *row* the style of *template_row*, so no leftover formatting remains."""
    if row == template_row:
        return
    for col in range(first_col, last_col + 1):
        target = worksheet._cells.get((row, col))
        source = worksheet._cells.get((template_row, col))
        if target is None or not is_single_cell(target):
            continue
        if source is not None and is_single_cell(source):
            target._style = copy(source._style)


def _totals_row_count(table: Table) -> int:
    """Number of totals rows in *table* (zero or one).

    Only ``totalsRowCount`` counts: ``totalsRowShown`` means "was shown at some
    point", not "is shown" (ECMA-376), matching COM's ``ShowTotals``.

    Examples:
        >>> from openpyxl.worksheet.table import Table, TableColumn
        >>> columns = [TableColumn(id=1, name="a"), TableColumn(id=2, name="b")]
        >>> _totals_row_count(Table(displayName="T", ref="A1:B2", tableColumns=columns))
        0
        >>> _totals_row_count(
        ...     Table(displayName="T", ref="A1:B3", tableColumns=columns, totalsRowShown=True)
        ... )
        0
        >>> _totals_row_count(
        ...     Table(displayName="T", ref="A1:B3", tableColumns=columns, totalsRowCount=1)
        ... )
        1

    """
    count = getattr(table, "totalsRowCount", None)
    return int(count) if count else 0


class _CellSnapshot(NamedTuple):
    """Value, style and type of a cell captured before the cell is cleared."""

    value: Any
    style: Any
    data_type: str | None


def _read_row(
    worksheet: Worksheet,
    row: int,
    first_col: int,
    last_col: int,
) -> list[_CellSnapshot]:
    """Capture the cells in *row* between *first_col* and *last_col*."""
    snapshots: list[_CellSnapshot] = []
    for col in range(first_col, last_col + 1):
        cell = worksheet._cells.get((row, col))
        if cell is None or not is_single_cell(cell):
            snapshots.append(_CellSnapshot(None, None, None))
            continue
        snapshots.append(_CellSnapshot(cell.value, copy(cell._style), cell.data_type))
    return snapshots


def _write_row(
    worksheet: Worksheet,
    row: int,
    first_col: int,
    snapshots: list[_CellSnapshot],
) -> None:
    """Write snapshots captured by :func:`_read_row` into *row*."""
    for offset, snapshot in enumerate(snapshots):
        cell = worksheet.cell(row=row, column=first_col + offset)
        if not is_single_cell(cell):
            continue
        if snapshot.style is not None:
            cell._style = copy(snapshot.style)
        cell.value = snapshot.value
        if snapshot.data_type is not None:
            cell.data_type = snapshot.data_type


def _extend_conditional_formatting(
    worksheet: Worksheet,
    table_start_column: int,
    table_end_column: int,
    first_data_row: int,
    table_last_row: int,
) -> None:
    """Extend conditional formatting rules from first data row to all rows in the table.

    A rule that tracks the table's height — one of its ranges starts on the
    first data row and lies inside the table's columns — is resized to the
    table's last row.  Ranges are parsed (never compared as strings), so a
    rule on ``B130:B200`` is not mistaken for one on ``B13``, and multi-area
    ranges such as ``B13:B20 D13:D20`` are handled per area.

    Args:
        worksheet: The worksheet containing the table
        table_start_column: 1-based index of the table's first column
        table_end_column: 1-based index of the table's last column
        first_data_row: Row number of the first data row
        table_last_row: Row number of the last data row in the table
            (the totals row, if present, is below it)

    """
    entries: list[Any] = list(worksheet.conditional_formatting)
    for conditional_formatting in entries:
        new_ranges: list[str] = []
        extended = False

        cell_ranges: list[Any] = sorted(conditional_formatting.sqref.ranges, key=str)
        for cell_range in cell_ranges:
            if _cf_range_tracks_table(
                cell_range,
                table_start_column=table_start_column,
                table_end_column=table_end_column,
                first_data_row=first_data_row,
                table_last_row=table_last_row,
            ):
                new_ranges.append(
                    CellRange(
                        min_col=cell_range.min_col,
                        min_row=cell_range.min_row,
                        max_col=cell_range.max_col,
                        max_row=table_last_row,
                    ).coord,
                )
                extended = True
            else:
                new_ranges.append(cell_range.coord)

        if not extended:
            continue

        # The rules are re-added under the new sqref instead of mutating
        # ``sqref`` in place: ConditionalFormattingList keys its entries by
        # the object, whose hash includes the range.
        old_sqref = str(conditional_formatting.sqref)
        new_sqref = " ".join(new_ranges)
        rules = list(conditional_formatting.rules)
        del worksheet.conditional_formatting[old_sqref]
        for rule in rules:
            worksheet.conditional_formatting.add(new_sqref, rule)

        logger.debug("Extended conditional formatting: %s -> %s", old_sqref, new_sqref)


def _cf_range_tracks_table(
    cell_range: CellRange,
    *,
    table_start_column: int,
    table_end_column: int,
    first_data_row: int,
    table_last_row: int,
) -> bool:
    """Whether *cell_range* should be resized to the table's last row.

    Such a range starts on the first data row and lies inside the table's
    columns, so it grows and shrinks with the table.

    Examples:
        >>> def tracks(area, last_row):
        ...     return _cf_range_tracks_table(
        ...         CellRange(area),
        ...         table_start_column=2,
        ...         table_end_column=13,
        ...         first_data_row=13,
        ...         table_last_row=last_row,
        ...     )
        >>> tracks("B13:B20", 30)
        True
        >>> tracks("B13:B20", 20)
        False
        >>> tracks("B130:B200", 30)
        False

    """
    return (
        cell_range.min_row == first_data_row
        and cell_range.max_row != table_last_row
        and cell_range.min_col >= table_start_column
        and cell_range.max_col <= table_end_column
    )
