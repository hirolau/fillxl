"""Chart integration tests — unique coverage not in story suite.

Tests chartEx warning emission, axis-position preservation, and exact
worksheet-name matching — behaviours that the story chart tests don't cover.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableColumn

from fillxl import fill
from tests.conftest import CHARTS_TEMPLATE_2

if TYPE_CHECKING:
    from pathlib import Path


def _build_two_sheet_chart_template(path: Path) -> None:
    """Build a workbook with tables (and charts) on ``Data`` and ``Data2``."""
    wb = openpyxl.Workbook()
    for sheet_title in ("Data", "Data2"):
        ws = wb.create_sheet(sheet_title) if sheet_title == "Data2" else wb.active
        ws.title = sheet_title
        ws["A1"] = "Category"
        ws["B1"] = "Value"
        for row in range(2, 4):
            ws.cell(row=row, column=1, value=f"cat{row - 1}")
            ws.cell(row=row, column=2, value=(row - 1) * 10)
        ws.add_table(
            Table(
                displayName=f"Table{sheet_title}",
                ref="A1:B3",
                tableColumns=[TableColumn(id=1, name="Category"), TableColumn(id=2, name="Value")],
            ),
        )

        chart = BarChart()
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
        charts_ws = wb.create_sheet(f"{sheet_title} Chart")
        charts_ws.add_chart(chart, "A1")

    wb.save(path)
    wb.close()


def _series_end_rows(path: Path) -> dict[str, set[str]]:
    """Map each chart's sheet name to the data references of its series."""
    wb = openpyxl.load_workbook(path)
    try:
        result: dict[str, set[str]] = {}
        for sheet_name in wb.sheetnames:
            for chart in getattr(wb[sheet_name], "_charts", []):
                refs = {
                    reference.f
                    for series in chart.ser
                    for label in (series.cat, series.val)
                    for reference in (
                        getattr(label, "numRef", None),
                        getattr(label, "strRef", None),
                    )
                    if reference is not None and reference.f
                }
                result[sheet_name] = refs
        return result
    finally:
        wb.close()


def test_a_chart_on_sheet_Data2_is_untouched_when_Data_is_filled(  # noqa: N802
    output_path: Path,
) -> None:
    """Sheet "Data" must not match sheet "Data2" by string prefix."""
    template = output_path.with_name("two_sheet_charts.xlsx")
    _build_two_sheet_chart_template(template)

    fill(
        template=template,
        output=output_path,
        tables={
            "TableData": [
                {"Category": "a", "Value": 1},
                {"Category": "b", "Value": 2},
                {"Category": "c", "Value": 3},
                {"Category": "d", "Value": 4},
            ],
        },
    )

    refs = _series_end_rows(output_path)
    assert refs["Data Chart"], refs
    assert refs["Data2 Chart"], refs

    # The chart bound to "Data" follows the table down to row 5 …
    assert all(ref.endswith("$5") for ref in refs["Data Chart"]), refs
    # … while the one bound to "Data2" still points at its original rows 2..3.
    assert all(ref.endswith("$3") for ref in refs["Data2 Chart"]), refs


@pytest.mark.filterwarnings("ignore:.*chartEx.*:UserWarning")
def test_fill_with_chartex_template_emits_warning(
    output_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """fill() should emit a warning when the template contains chartEx charts."""
    caplog.set_level("WARNING")

    fill(
        template=CHARTS_TEMPLATE_2,
        output=output_path,
        tables={"WaterfallChartTable": [{"Category": "A", "Amount": 100}]},
    )

    assert any(
        "chartEx" in record.message or "waterfall" in record.message
        for record in caplog.records
        if record.levelname == "WARNING"
    ), f"Logger warning not found. Captured: {[r.message for r in caplog.records]}"


@pytest.mark.filterwarnings("ignore:.*chartEx.*:UserWarning")
def test_area_chart_category_axis_at_bottom_after_fill(output_path: Path) -> None:
    """Area chart catAx should have axPos='b' (bottom) after filling.

    Tests both that:
    1. The category axis is correctly positioned (axPos='b')
    2. The cache ptCount matches the actual data range (no stale cache)
    """
    import re
    import zipfile

    fill(
        template=CHARTS_TEMPLATE_2,
        output=output_path,
        tables={
            "AreaChartTable": [
                {"Week": "Mon", "Desktop": 100, "Mobile": 200, "Tablet": 30},
                {"Week": "Tue", "Desktop": 150, "Mobile": 250, "Tablet": 40},
                {"Week": "Wed", "Desktop": 200, "Mobile": 300, "Tablet": 50},
            ],
        },
    )

    with zipfile.ZipFile(output_path) as z:
        chart_xml = z.read("xl/charts/chart5.xml").decode("utf-8")

        # ── 1. Verify catAx position ──
        cat_blk = re.search(r"<[^ >]*:?catAx[ >].*?</[^ >]*:?catAx>", chart_xml, re.DOTALL)
        assert cat_blk is not None, "No catAx block found in Area chart XML"
        cat_xml = cat_blk.group()

        pos = re.search(r'axPos val="(\w+)"', cat_xml)
        lbl = re.search(r'tickLblPos val="(\w+)"', cat_xml)

        assert pos is not None, "No axPos found in catAx"
        assert pos.group(1) == "b", (
            f"Area chart category axis should be at bottom (axPos='b'), got axPos={pos.group(1)!r}"
        )
        assert lbl is not None, "No tickLblPos found in catAx"
        assert lbl.group(1) == "nextTo", (
            f"Area chart tick labels should be visible (tickLblPos='nextTo'), "
            f"got tickLblPos={lbl.group(1)!r}"
        )

        # ── 2. Verify cache ptCount matches data range ──
        for ref_type in ("strRef", "numRef"):
            for blk in re.findall(rf"<{ref_type}[ >].*?</{ref_type}>", chart_xml, re.DOTALL):
                f_match = re.search(r"<[^ >]*:?f[ >]([^<]+)", blk)
                if not f_match:
                    continue
                f_text = f_match.group(1)
                if ":" not in f_text.split("!")[-1]:
                    continue

                pc_match = re.search(r'ptCount val="(\d+)"', blk)
                assert pc_match is not None, f"No ptCount found in {ref_type} for {f_text}"
                ref_range = f_text.split("!")[-1]
                ref_parts = ref_range.split(":")
                start_cell = ref_parts[0]
                end_cell = ref_parts[1]
                start_row = int("".join(ch for ch in start_cell if ch.isdigit()))
                end_row = int("".join(ch for ch in end_cell if ch.isdigit()))
                expected_pt_count = end_row - start_row + 1
                actual_pt_count = int(pc_match.group(1))

                assert actual_pt_count == expected_pt_count, (
                    f"ptCount mismatch in {ref_type} for {f_text}: "
                    f"expected {expected_pt_count}, got {actual_pt_count}"
                )

        # ── 3. Verify critical catAx elements that openpyxl drops ──
        for elem, expected in [
            ("delete", "0"),
            ("crosses", "autoZero"),
            ("auto", "1"),
            ("lblAlgn", "ctr"),
            ("noMultiLvlLbl", "0"),
        ]:
            m = re.search(elem + r' val="([^"]+)"', cat_xml)
            assert m is not None, f"catAx missing <{elem}> element"
            assert m.group(1) == expected, (
                f"catAx {elem} should be {expected!r}, got {m.group(1)!r}"
            )

        # ── 4. Verify critical valAx elements ──
        val_blk = re.search(r"<[^ >]*:?valAx[ >].*?</[^ >]*:?valAx>", chart_xml, re.DOTALL)
        assert val_blk is not None, "No valAx block found in Area chart XML"
        val_xml = val_blk.group()

        for elem, expected in [
            ("delete", "0"),
            ("crosses", "autoZero"),
            ("crossBetween", "midCat"),
        ]:
            m = re.search(elem + r' val="([^"]+)"', val_xml)
            assert m is not None, f"valAx missing <{elem}> element"
            assert m.group(1) == expected, (
                f"valAx {elem} should be {expected!r}, got {m.group(1)!r}"
            )
