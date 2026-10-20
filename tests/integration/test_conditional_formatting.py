"""Conditional formatting integration tests — unique coverage not in story suite.

Tests behaviours not covered by story tests: disabled CF flag, COM-based
DisplayFormat verification, and multi-column CF rule preservation.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest
from openpyxl.worksheet.table import Table, TableColumn

from fillxl import FillXLConfig, fill
from tests._assets import TEMPLATE

if TYPE_CHECKING:
    from pathlib import Path

    import xlwings as xw


def _data_rows(count: int) -> list[dict[str, object]]:
    """Generate *count* rows of table data."""
    return [
        {
            "Col1": f"row{i}",
            "Col2": i,
            "ColNotFilled": f"nf{i}",
            "ColFormula": None,
            "TestFormattedPct": 0.0,
        }
        for i in range(1, count + 1)
    ]


def _get_cf_sqrefs(output_path: Path) -> set[str]:
    """Return the set of CF sqref strings from the output workbook."""
    wb = openpyxl.load_workbook(output_path)
    try:
        ws = wb["TestTable"]
        return {str(cf.sqref) for cf in ws.conditional_formatting}
    finally:
        wb.close()


def test_conditional_formatting_remains_unchanged_when_extend_flag_is_false(
    output_path: Path,
) -> None:
    """extend_conditional_formatting=False leaves CF ranges unchanged."""
    fill(
        template=TEMPLATE,
        output=output_path,
        tables={"TableTest": _data_rows(10)},
        config=FillXLConfig(extend_conditional_formatting=False),
    )

    sqrefs = _get_cf_sqrefs(output_path)
    assert any("E5" in s for s in sqrefs)
    assert not any("E11" in s for s in sqrefs)


@pytest.mark.xlwings
def test_conditional_formatting_applies_green_background_for_value_equal_to_seven(
    output_path: Path,
    xl_app: xw.App,
) -> None:
    """A cell with value 7 should have a green background via DisplayFormat."""
    data = [
        {
            "Col1": "row1",
            "Col2": 1,
            "ColNotFilled": "x",
            "ColFormula": None,
            "TestFormattedPct": 0.0,
        },
        {"Col1": 7, "Col2": 2, "ColNotFilled": "y", "ColFormula": None, "TestFormattedPct": 0.0},
    ]

    fill(
        template=TEMPLATE,
        output=output_path,
        tables={"TableTest": data},
    )

    wb = xl_app.books.open(str(output_path))
    try:
        ws = wb.sheets["TestTable"]
        cell_seven = ws.range("A3")
        cell_normal = ws.range("B2")
        green_color = cell_seven.api.DisplayFormat.Interior.Color
        normal_color = cell_normal.api.DisplayFormat.Interior.Color
        assert green_color != normal_color, (
            f"Expected green CF fill ({green_color}) to differ from normal fill ({normal_color})"
        )
        assert green_color != 16247773, (
            f"Expected 7-cell to have CF-applied green fill, got default {green_color}"
        )
    finally:
        wb.close()


@pytest.mark.xlwings
def test_conditional_formatting_applies_red_background_for_value_between_two_and_three(
    output_path: Path,
    xl_app: xw.App,
) -> None:
    """Col2 cell with value 2.5 should have a red background via DisplayFormat."""
    data = [
        {"Col1": "a", "Col2": 1, "ColNotFilled": "x", "ColFormula": None, "TestFormattedPct": 0.0},
        {
            "Col1": "b",
            "Col2": 2.5,
            "ColNotFilled": "y",
            "ColFormula": None,
            "TestFormattedPct": 0.0,
        },
    ]

    fill(
        template=TEMPLATE,
        output=output_path,
        tables={"TableTest": data},
    )

    wb = xl_app.books.open(str(output_path))
    try:
        ws = wb.sheets["TestTable"]
        cell_red = ws.range("B3")
        cell_normal = ws.range("B2")
        red_color = cell_red.api.DisplayFormat.Interior.Color
        normal_color = cell_normal.api.DisplayFormat.Interior.Color
        assert red_color == 255, f"Expected red CF fill (255) for Col2=2.5, got {red_color}"
        assert normal_color != 255, f"Expected normal cell to not be red, got {normal_color}"
    finally:
        wb.close()


def test_conditional_formatting_preserves_multi_column_rules_during_extension(
    output_path: Path,
) -> None:
    """Multi-column CF rules should not be corrupted by extension."""
    fill(
        template=TEMPLATE,
        output=output_path,
        tables={"TableTest": _data_rows(8)},
    )

    wb = openpyxl.load_workbook(output_path)
    try:
        ws = wb["TestTable"]
        cf_count_original = len(
            {
                str(cf.sqref)
                for cf in openpyxl.load_workbook(TEMPLATE)["TestTable"].conditional_formatting
            },
        )
        cf_count_after = len({str(cf.sqref) for cf in ws.conditional_formatting})
        assert cf_count_after >= cf_count_original, (
            f"CF rules dropped: {cf_count_original} → {cf_count_after}"
        )
        for cf in ws.conditional_formatting:
            sqref = str(cf.sqref)
            end_row = sqref.rsplit(":", maxsplit=1)[-1] if ":" in sqref else sqref
            end_row_num = int("".join(ch for ch in end_row if ch.isdigit()))
            assert end_row_num <= 9, f"CF sqref {sqref} extends beyond row 9"
    finally:
        wb.close()


# ═══════════════════════════════════════════════════════════════════════════
# Range parsing: rules are matched by geometry, never by string prefix
# ═══════════════════════════════════════════════════════════════════════════


def _build_row_12_table_template(path: Path, *, extra_rule_range: str) -> None:
    """Build a 3-column table whose data rows start at row 13.

    Row 13 is what makes a rule on row 130 a realistic near-miss for
    prefix-based matching.

    Args:
        path: Where to save the template.
        extra_rule_range: Second CF rule range to add.

    """
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for col_idx, name in enumerate(["Category", "Value", "Other"], 1):
        ws.cell(row=12, column=col_idx, value=name)
    for row in (13, 14):
        ws.cell(row=row, column=1, value=f"sample{row}")
        ws.cell(row=row, column=2, value=row)
        ws.cell(row=row, column=3, value=row)

    ws.conditional_formatting.add(
        "B13:B14",
        CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill()),
    )
    ws.conditional_formatting.add(
        extra_rule_range,
        CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill()),
    )

    ws.add_table(
        Table(
            displayName="RowTwelve",
            ref="A12:C14",
            tableColumns=[
                TableColumn(id=1, name="Category"),
                TableColumn(id=2, name="Value"),
                TableColumn(id=3, name="Other"),
            ],
        ),
    )
    wb.save(path)
    wb.close()


def _cf_sqref_text(path: Path) -> str:
    """All CF sqref strings of the Data sheet, joined for containment checks."""
    wb = openpyxl.load_workbook(path)
    try:
        return " ".join(sorted(str(cf.sqref) for cf in wb["Data"].conditional_formatting))
    finally:
        wb.close()


def test_conditional_formatting_on_B130_is_not_mistaken_for_B13(  # noqa: N802
    output_path: Path,
) -> None:
    """A rule on B130:B200 must survive a fill of a table starting at B13."""
    template = output_path.with_name("cf_b130_template.xlsx")
    _build_row_12_table_template(template, extra_rule_range="B130:B200")

    fill(
        template=template,
        output=output_path,
        tables={
            "RowTwelve": [{"Category": f"c{i}", "Value": i, "Other": i} for i in range(1, 6)],
        },
    )

    sqref_text = _cf_sqref_text(output_path)
    assert "B13:B17" in sqref_text, sqref_text
    assert "B130:B200" in sqref_text, sqref_text
    assert "B130:B17" not in sqref_text, sqref_text


def test_a_multi_area_conditional_formatting_range_is_resized_per_area(
    output_path: Path,
) -> None:
    """Every area of a multi-area sqref is resized independently."""
    template = output_path.with_name("cf_multi_area_template.xlsx")
    _build_row_12_table_template(template, extra_rule_range="B13 C13")

    fill(
        template=template,
        output=output_path,
        tables={
            "RowTwelve": [{"Category": f"c{i}", "Value": i, "Other": i} for i in range(1, 6)],
        },
    )

    sqref_text = _cf_sqref_text(output_path)
    assert "B13:B17" in sqref_text, sqref_text
    assert "C13:C17" in sqref_text, sqref_text


def test_conditional_formatting_below_the_table_is_left_alone(
    output_path: Path,
) -> None:
    """A rule that does not start on the first data row is not resized."""
    template = output_path.with_name("cf_below_template.xlsx")
    _build_row_12_table_template(template, extra_rule_range="C20:C40")

    fill(
        template=template,
        output=output_path,
        tables={
            "RowTwelve": [{"Category": f"c{i}", "Value": i, "Other": i} for i in range(1, 6)],
        },
    )

    sqref_text = _cf_sqref_text(output_path)
    assert "C20:C40" in sqref_text, sqref_text
    assert "B13:B17" in sqref_text, sqref_text
