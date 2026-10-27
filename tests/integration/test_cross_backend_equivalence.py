"""Cross-backend equivalence.

The promise of fillxl is that the openpyxl and win32com engines produce the
same workbook for the same input.  Every test here fills one template with
both engines, reads the results back with openpyxl and compares them.

The win32com cases skip themselves when Excel is not available (the Linux CI
runners), so the same suite protects local runs and CI pipelines with Excel.
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from typing import TYPE_CHECKING, Any

import numpy as np
import openpyxl
import pandas as pd
import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableColumn, TableFormula

from fillxl import CellValueError, Engine, FillXLConfig, FillXLWarning, fill
from tests._assets import TEMPLATE
from tests._engines import ENGINES, _fill_output, _win32com_available

if TYPE_CHECKING:
    from pathlib import Path

COLUMNS = ["Name", "Value", "Notes"]

# One value per interesting dtype.  Both engines must agree on all of them.
VALUES = [
    pytest.param(42, id="int"),
    pytest.param(np.int64(42), id="numpy_int64"),
    pytest.param(np.float32(1.5), id="numpy_float32"),
    pytest.param(Decimal("12.34"), id="decimal"),
    pytest.param(dt.date(2025, 12, 25), id="date"),
    pytest.param(dt.time(9, 15, 30), id="time"),
    pytest.param(dt.datetime(2025, 6, 15, 14, 30), id="datetime"),
    pytest.param(
        dt.datetime(2025, 6, 15, 14, 30, tzinfo=dt.timezone(dt.timedelta(hours=2))),
        id="tz_datetime",
    ),
    pytest.param(np.datetime64("2025-06-15T14:30"), id="numpy_datetime64"),
    pytest.param(dt.timedelta(hours=6), id="timedelta"),
    pytest.param(float("nan"), id="nan"),
    pytest.param(True, id="bool"),
    pytest.param(None, id="none"),
    pytest.param("hello", id="text"),
    pytest.param("=1+1", id="formula_string"),
    pytest.param("00123", id="numeric_string"),
    pytest.param("1/2", id="slash_string"),
    pytest.param("#N/A", id="error_string"),
    pytest.param("", id="empty_string"),
]


@pytest.fixture(scope="module")
def equivalence_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """A template with one named cell and a one-row, three-column table."""
    from openpyxl.worksheet.filters import AutoFilter

    path = tmp_path_factory.mktemp("equivalence") / "equivalence.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for col_idx, name in enumerate(COLUMNS, 1):
        ws.cell(row=1, column=col_idx, value=name)
    ws.cell(row=2, column=1, value="sample")
    ws.cell(row=2, column=2, value=0)
    ws.cell(row=2, column=3, value="template note")
    wb.defined_names.add(DefinedName("probe_cell", attr_text="Data!$E$1"))
    wb.defined_names.add(DefinedName("probe_block", attr_text="Data!$E$2:$F$3"))
    ws.add_table(
        Table(
            displayName="Items",
            ref="A1:C2",
            tableColumns=[TableColumn(id=i + 1, name=name) for i, name in enumerate(COLUMNS)],
            autoFilter=AutoFilter(ref="A1:C2"),
        ),
    )
    wb.save(path)
    wb.close()
    return path


class FilledWorkbook:
    """The read-back of one filled workbook."""

    def __init__(self, worksheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
        self.worksheet = worksheet
        self.named_cell = worksheet["E1"].value
        self.first_row = [worksheet.cell(row=2, column=col).value for col in range(1, 4)]
        self.table = next(iter(worksheet.tables.values()))


def _fill_and_read(
    template: Path,
    output_path: Path,
    engine: Engine,
    *,
    values: dict[str, Any] | None = None,
    tables: dict[str, Any] | None = None,
    config: FillXLConfig | None = None,
) -> FilledWorkbook:
    """Fill *template* with *engine* and return the openpyxl read-back."""
    try:
        output = _fill_output(
            template,
            output_path,
            engine,
            values=values,
            tables=tables,
            config=config,
        )
    except Exception as exc:
        message = str(exc).lower()
        if engine == "win32com" and ("win32com" in message or "excel" in message):
            pytest.skip(f"win32com backend not available: {exc}")
        raise

    wb = openpyxl.load_workbook(output)
    try:
        return FilledWorkbook(wb["Data"])
    finally:
        wb.close()


def _table_rows(value: object) -> list[dict[str, object]]:
    """One table row carrying *value* in the Value column."""
    return [{"Name": "row", "Value": value, "Notes": "note"}]


# ═══════════════════════════════════════════════════════════════════════════════
# Same value, both engines
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("value", VALUES)
def test_both_engines_store_the_same_named_cell_value(
    equivalence_template: Path,
    output_path: Path,
    value: object,
) -> None:
    """A named cell receives the same value from both engines."""
    opened = _fill_and_read(
        equivalence_template,
        output_path,
        "openpyxl",
        values={"probe_cell": value},
    )
    closed = _fill_and_read(
        equivalence_template,
        output_path,
        "win32com",
        values={"probe_cell": value},
    )

    if isinstance(value, float) and value != value:  # NaN
        assert opened.named_cell is None
        assert closed.named_cell is None
    elif isinstance(value, Decimal):
        assert isinstance(opened.named_cell, (int, float))
        assert isinstance(closed.named_cell, (int, float))
        assert float(opened.named_cell) == pytest.approx(float(closed.named_cell))
    else:
        assert opened.named_cell == closed.named_cell


@pytest.mark.parametrize("value", VALUES)
def test_both_engines_store_the_same_table_value(
    equivalence_template: Path,
    output_path: Path,
    value: object,
) -> None:
    """A table cell receives the same value from both engines."""
    opened = _fill_and_read(
        equivalence_template,
        output_path,
        "openpyxl",
        tables={"Items": _table_rows(value)},
    )
    closed = _fill_and_read(
        equivalence_template,
        output_path,
        "win32com",
        tables={"Items": _table_rows(value)},
    )

    if isinstance(value, float) and value != value:  # NaN
        assert opened.first_row[1] is None
        assert closed.first_row[1] is None
    elif isinstance(value, Decimal):
        assert isinstance(opened.first_row[1], (int, float))
        assert isinstance(closed.first_row[1], (int, float))
        assert float(opened.first_row[1]) == pytest.approx(float(closed.first_row[1]))
    else:
        assert opened.first_row[1] == closed.first_row[1]


def test_both_engines_write_the_same_datetime_to_a_named_cell(
    equivalence_template: Path,
    output_path: Path,
) -> None:
    """A named-cell datetime is not shifted by the local timezone offset.

    win32com converts naive datetimes with the local offset, so the value has
    to be marked as UTC before it is written — otherwise the two engines store
    timestamps hours apart.
    """
    moment = dt.datetime(2025, 6, 15, 14, 30, 0)

    opened = _fill_and_read(
        equivalence_template,
        output_path,
        "openpyxl",
        values={"probe_cell": moment},
    )
    closed = _fill_and_read(
        equivalence_template,
        output_path,
        "win32com",
        values={"probe_cell": moment},
    )

    assert opened.named_cell == moment
    assert closed.named_cell == moment


@pytest.mark.parametrize("engine", ENGINES)
def test_numpy_integers_land_in_excel_as_numbers_not_text(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """np.int64 values must not be written as strings."""
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        tables={"Items": [{"Name": "a", "Value": np.int64(7), "Notes": "n"}]},
    )

    value = filled.first_row[1]
    assert value == 7
    assert not isinstance(value, str)
    assert filled.worksheet["B2"].data_type == "n"


# ═══════════════════════════════════════════════════════════════════════════════
# Tables: shape, autoFilter, empty data, missing columns
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_a_resized_table_keeps_its_autofilter_in_step(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """table.ref and table.autoFilter.ref must stay identical.

    A mismatch is a classic cause of Excel's "we found a problem with some
    content" repair prompt.
    """
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        tables={"Items": [{"Name": f"row{i}", "Value": i, "Notes": "n"} for i in range(3)]},
    )

    assert filled.table.ref == "A1:C4"
    assert filled.table.autoFilter.ref == filled.table.ref


@pytest.mark.parametrize("engine", ENGINES)
def test_an_empty_dataframe_leaves_a_valid_one_row_table(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """An Excel table needs one body row, so empty data leaves one cleared row."""
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        tables={"Items": pd.DataFrame(columns=COLUMNS)},
    )

    assert filled.table.ref == "A1:C2"
    assert filled.table.autoFilter.ref == "A1:C2"
    assert filled.first_row == [None, None, None]


@pytest.mark.parametrize("engine", ENGINES)
def test_a_column_missing_from_the_data_does_not_leak_template_values(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """Sample values in the template are cleared for columns the data omits."""
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        tables={"Items": [{"Name": f"row{i}", "Value": i} for i in range(3)]},
        config=FillXLConfig(column_validation_strategy="warn"),
    )

    assert [filled.worksheet.cell(row=row, column=3).value for row in (2, 3, 4)] == [
        None,
        None,
        None,
    ]


# ═══════════════════════════════════════════════════════════════════════════════
# Styles and conditional formatting
# ═══════════════════════════════════════════════════════════════════════════════


def _cf_ranges(path: Path, sheet: str) -> set[str]:
    wb = openpyxl.load_workbook(path)
    try:
        return {str(cf.sqref) for cf in wb[sheet].conditional_formatting}
    finally:
        wb.close()


def _cf_rule_count(path: Path, sheet: str) -> int:
    wb = openpyxl.load_workbook(path)
    try:
        return sum(len(list(cf.rules)) for cf in wb[sheet].conditional_formatting)
    finally:
        wb.close()


def test_both_engines_extend_conditional_formatting_to_the_same_ranges(
    output_path: Path,
) -> None:
    """The COM path uses the same range rule as the openpyxl path."""
    rows = [
        {"Col1": f"row{i}", "Col2": i, "ColNotFilled": "n", "TestFormattedPct": 0.1}
        for i in range(5)
    ]
    opened = _fill_output(
        TEMPLATE,
        output_path.with_name("cf_openpyxl.xlsx"),
        "openpyxl",
        tables={"TableTest": rows},
    )
    closed = _fill_output(
        TEMPLATE,
        output_path.with_name("cf_win32com.xlsx"),
        "win32com",
        tables={"TableTest": rows},
    )

    assert _cf_ranges(closed, "TestTable") == _cf_ranges(opened, "TestTable")


@pytest.mark.parametrize("engine", ENGINES)
def test_conditional_formatting_rule_count_is_unchanged_after_fill(
    output_path: Path,
    engine: Engine,
) -> None:
    """Extending a range must not fragment or duplicate rules."""
    rows = [
        {"Col1": f"row{i}", "Col2": i, "ColNotFilled": "n", "TestFormattedPct": 0.1}
        for i in range(5)
    ]
    output = _fill_output(TEMPLATE, output_path, engine, tables={"TableTest": rows})

    assert _cf_rule_count(output, "TestTable") == _cf_rule_count(TEMPLATE, "TestTable")


@pytest.mark.parametrize("engine", ENGINES)
def test_new_rows_get_the_first_row_number_format(
    output_path: Path,
    engine: Engine,
) -> None:
    """New rows copy the first data row's number formats."""
    rows = [
        {"Col1": f"row{i}", "Col2": i, "ColNotFilled": "n", "TestFormattedPct": 0.1}
        for i in range(6)
    ]
    output = _fill_output(
        TEMPLATE,
        output_path,
        engine,
        tables={"TableTest": rows},
        config=FillXLConfig(copy_table_styles=True),
    )

    wb = openpyxl.load_workbook(output)
    try:
        ws = wb["TestTable"]
        assert ws["E7"].number_format == ws["E2"].number_format
        assert ws["E2"].number_format != "General"
    finally:
        wb.close()


@pytest.mark.parametrize(
    "engine",
    [
        pytest.param("openpyxl", id="openpyxl"),
        pytest.param(
            "win32com",
            id="win32com",
            marks=pytest.mark.xfail(
                strict=True,
                reason="Excel applies the table column format to new rows itself",
            ),
        ),
    ],
)
def test_copy_table_styles_false_does_not_copy_number_formats(
    output_path: Path,
    engine: Engine,
) -> None:
    """Without style copying the new rows keep their own (default) format."""
    rows = [
        {"Col1": f"row{i}", "Col2": i, "ColNotFilled": "n", "TestFormattedPct": 0.1}
        for i in range(6)
    ]
    output = _fill_output(
        TEMPLATE,
        output_path,
        engine,
        tables={"TableTest": rows},
        config=FillXLConfig(copy_table_styles=False),
    )

    wb = openpyxl.load_workbook(output)
    try:
        assert wb["TestTable"]["E7"].number_format == "General"
    finally:
        wb.close()


@pytest.mark.parametrize("engine", ENGINES)
def test_a_mixed_type_column_keeps_number_formats_per_cell(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A date in the middle of a numeric column formats only its own cell."""
    output = _fill_output(
        equivalence_template,
        output_path,
        engine,
        tables={
            "Items": [
                {"Name": "a", "Value": 1, "Notes": "n"},
                {"Name": "b", "Value": dt.datetime(2025, 1, 1), "Notes": "n"},
                {"Name": "c", "Value": 3, "Notes": "n"},
            ]
        },
    )

    wb = openpyxl.load_workbook(output)
    try:
        ws = wb["Data"]
        assert ws["B2"].number_format == "General"
        assert ws["B4"].number_format == "General"
        assert ws["B3"].number_format != "General"
    finally:
        wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# Formula escaping
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_formula_like_text_stays_literal_by_default(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A string starting with "=" is literal text in both engines."""
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        values={"probe_cell": "=1+1"},
        tables={"Items": _table_rows("=1+1")},
    )

    assert filled.named_cell == "=1+1"
    assert filled.worksheet["E1"].data_type == "s"
    assert filled.first_row[1] == "=1+1"
    assert filled.worksheet["B2"].data_type == "s"


@pytest.mark.parametrize("engine", ENGINES)
def test_escape_formulas_false_writes_formulas_in_both_engines(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """With escape_formulas=False the "=" string becomes a formula."""
    config = FillXLConfig(escape_formulas=False)
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        values={"probe_cell": "=1+1"},
        tables={"Items": _table_rows("=1+1")},
        config=config,
    )

    assert filled.named_cell == "=1+1"
    assert filled.worksheet["E1"].data_type == "f"
    assert filled.first_row[1] == "=1+1"
    assert filled.worksheet["B2"].data_type == "f"


@pytest.mark.parametrize("engine", ENGINES)
def test_numeric_looking_strings_stay_text_in_both_engines(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """COM parses "00123" as typed input; the apostrophe prefix prevents it."""
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        values={"probe_cell": "00123"},
        tables={"Items": _table_rows("00123")},
    )

    assert filled.named_cell == "00123"
    assert filled.first_row[1] == "00123"
    assert filled.worksheet["B2"].data_type == "s"


@pytest.mark.parametrize("engine", ENGINES)
def test_error_looking_strings_are_stored_as_text_not_errors(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """An error-looking string is data, not an Excel error, in both engines.

    openpyxl turns strings that match an error code into error cells on its
    own, so the cell type has to be forced to text.
    """
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        values={"probe_cell": "#N/A"},
        tables={"Items": _table_rows("#N/A")},
    )

    assert filled.named_cell == "#N/A"
    assert filled.worksheet["E1"].data_type == "s"
    assert filled.first_row[1] == "#N/A"
    assert filled.worksheet["B2"].data_type == "s"


# ═══════════════════════════════════════════════════════════════════════════════
# Excel cell limits
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_overlong_text_raises_with_table_row_and_column(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A cell holds at most 32,767 characters."""
    rows = [{"Name": "a", "Value": 1, "Notes": "x" * 32_768}]

    with pytest.raises(CellValueError, match=r"Table 'Items', data row 1, column 'Notes'"):
        _fill_output(equivalence_template, output_path, engine, tables={"Items": rows})


@pytest.mark.parametrize("engine", ENGINES)
def test_a_control_character_raises_with_context(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """Excel rejects control characters in text."""
    rows = [{"Name": "a", "Value": 1, "Notes": "bell\x07"}]

    with pytest.raises(CellValueError, match="control character"):
        _fill_output(equivalence_template, output_path, engine, tables={"Items": rows})


@pytest.mark.parametrize("engine", ENGINES)
def test_a_date_before_march_1900_raises(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """Excel counts a 29 February 1900, so earlier dates are unreliable."""
    rows = [{"Name": "a", "Value": dt.date(1900, 2, 28), "Notes": "n"}]

    with pytest.raises(CellValueError, match="before 1900-03-01"):
        _fill_output(equivalence_template, output_path, engine, tables={"Items": rows})


@pytest.mark.parametrize("engine", ENGINES)
def test_a_large_integer_warns_once_per_column(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
    recwarn: pytest.WarningsRecorder,
) -> None:
    """Excel rounds integers above 15 significant digits; say so once."""
    rows = [{"Name": f"r{i}", "Value": 1234567890123456, "Notes": "n"} for i in range(3)]

    _fill_output(equivalence_template, output_path, engine, tables={"Items": rows})

    precision = [w for w in recwarn if issubclass(w.category, FillXLWarning)]
    assert len(precision) == 1
    assert "significant digits" in str(precision[0].message)


@pytest.mark.parametrize("engine", ENGINES)
def test_a_named_cell_error_names_the_cell(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A named-cell error says which name failed."""
    with pytest.raises(CellValueError, match="Named cell 'probe_cell'"):
        _fill_output(
            equivalence_template,
            output_path,
            engine,
            values={"probe_cell": "x" * 32_768},
        )


# ═══════════════════════════════════════════════════════════════════════════════
# Duplicate column names
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_duplicate_normalized_column_names_are_rejected_by_both_engines(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """The adapter rejects the ambiguity before either engine maps names."""
    with pytest.raises(ValueError, match=r"Duplicate column names.*'1'"):
        _fill_output(
            equivalence_template,
            output_path,
            engine,
            tables={"Items": {1: ["a"], "1": ["b"]}},
        )


# ═══════════════════════════════════════════════════════════════════════════════
# Atomic output
# ═══════════════════════════════════════════════════════════════════════════════


def _engine_output(output_path: Path, engine: Engine) -> Path:
    return output_path.with_name(f"{output_path.stem}_{engine}{output_path.suffix}")


@pytest.mark.parametrize("engine", ENGINES)
def test_a_failed_fill_leaves_the_previous_output_untouched(
    output_path: Path,
    engine: Engine,
) -> None:
    """A failure must never replace an existing output."""
    target = _engine_output(output_path, engine)
    target.write_bytes(b"old")

    with pytest.raises(KeyError):
        _fill_output(TEMPLATE, output_path, engine, tables={"NoSuchTable": []})

    assert target.read_bytes() == b"old"


@pytest.mark.parametrize("engine", ENGINES)
def test_a_failed_fill_leaves_no_staging_file(output_path: Path, engine: Engine) -> None:
    """The staging file is removed when the fill fails."""
    with pytest.raises(KeyError):
        _fill_output(TEMPLATE, output_path, engine, tables={"NoSuchTable": []})

    assert list(output_path.parent.glob(".*.fillxl*")) == []


@pytest.mark.parametrize("engine", ENGINES)
def test_a_successful_fill_leaves_no_staging_file(output_path: Path, engine: Engine) -> None:
    """The staging file is renamed into place."""
    _fill_output(TEMPLATE, output_path, engine, values={"single_cell": "x"})

    assert list(output_path.parent.glob(".*.fillxl*")) == []


# ═══════════════════════════════════════════════════════════════════════════════
# Named cells
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_a_multi_cell_named_range_is_rejected_by_both_engines(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A scalar cannot be written to a multi-cell name on either engine."""
    with pytest.raises(ValueError, match="multi-cell range"):
        _fill_output(
            equivalence_template,
            output_path,
            engine,
            values={"probe_block": "x"},
        )


@pytest.mark.parametrize("engine", ENGINES)
def test_a_multi_destination_name_updates_only_the_first_cell(
    multi_dest_name_template: Path,
    output_path: Path,
    engine: Engine,
    recwarn: pytest.WarningsRecorder,
) -> None:
    """Both engines write to the first destination and point out the rest."""
    output = _fill_output(
        multi_dest_name_template,
        output_path,
        engine,
        values={"multi_dest": "new"},
    )

    wb = openpyxl.load_workbook(output)
    try:
        assert wb["Sheet1"]["A1"].value == "new"
        assert wb["Sheet2"]["B2"].value == "Second dest"
    finally:
        wb.close()
    ignored = [w for w in recwarn if issubclass(w.category, FillXLWarning)]
    assert any("destinations" in str(w.message) for w in ignored)


# ═══════════════════════════════════════════════════════════════════════════════
# Ragged list[dict] data
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_list_of_dicts_with_keys_first_seen_in_later_rows(
    equivalence_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A late key is written instead of being silently dropped."""
    filled = _fill_and_read(
        equivalence_template,
        output_path,
        engine,
        tables={
            "Items": [
                {"Name": "a", "Value": 1},
                {"Name": "b", "Value": 2, "Notes": "late"},
            ]
        },
    )

    assert filled.worksheet["C2"].value is None
    assert filled.worksheet["C3"].value == "late"


# ═══════════════════════════════════════════════════════════════════════════════
# Calculated columns
# ═══════════════════════════════════════════════════════════════════════════════


def _build_priced_template(path: Path, body_value: object) -> None:
    """A one-data-row table whose Price column has the definition ``100``."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Price"])
    ws.append(["old", body_value])
    ws.add_table(
        Table(
            displayName="Priced",
            ref="A1:B2",
            tableColumns=[
                TableColumn(id=1, name="Name"),
                TableColumn(
                    id=2,
                    name="Price",
                    calculatedColumnFormula=TableFormula(attr_text="100"),
                ),
            ],
        ),
    )
    wb.save(path)
    wb.close()


@pytest.fixture(scope="module")
def stale_definition_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """A table whose calculated-column definition is stale (the body holds a value)."""
    path = tmp_path_factory.mktemp("stale_definition") / "stale_definition.xlsx"
    _build_priced_template(path, body_value=5)
    return path


@pytest.fixture(scope="module")
def constant_formula_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """A table whose ``=100`` calculated column holds the formula in the body."""
    path = tmp_path_factory.mktemp("constant_formula") / "constant_formula.xlsx"
    _build_priced_template(path, body_value="=100")
    return path


@pytest.mark.parametrize("engine", ENGINES)
def test_a_calculated_column_survives_data_that_names_it(
    calculated_column_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """Data that names a formula column (even as None) must not erase the formula."""
    rows = [
        {
            "Col1": "a",
            "Col2": 1,
            "ColNotFilled": "x",
            "ColFormula": None,
            "TestFormattedPct": 0.1,
        },
        {
            "Col1": "b",
            "Col2": 2,
            "ColNotFilled": "y",
            "ColFormula": None,
            "TestFormattedPct": 0.2,
        },
    ]
    output = _fill_output(
        calculated_column_template, output_path, engine, tables={"TableTest": rows}
    )

    wb = openpyxl.load_workbook(output)
    try:
        ws = wb["TestTable"]
        for row in (2, 3):
            cell = ws.cell(row=row, column=4)
            assert cell.data_type == "f"
            assert str(cell.value).startswith("=")
    finally:
        wb.close()


@pytest.mark.parametrize("engine", ENGINES)
def test_real_data_for_a_calculated_column_is_ignored_with_a_warning(
    calculated_column_template: Path,
    output_path: Path,
    engine: Engine,
    recwarn: pytest.WarningsRecorder,
) -> None:
    """Real data for a formula column is dropped, and the user is told."""
    rows = [
        {
            "Col1": "a",
            "Col2": 1,
            "ColNotFilled": "x",
            "ColFormula": "oops",
            "TestFormattedPct": 0.1,
        },
    ]

    output = _fill_output(
        calculated_column_template, output_path, engine, tables={"TableTest": rows}
    )

    wb = openpyxl.load_workbook(output)
    try:
        cell = wb["TestTable"].cell(row=2, column=4)
        assert cell.data_type == "f"
        assert str(cell.value).startswith("=")
    finally:
        wb.close()
    ignored = [w for w in recwarn if issubclass(w.category, FillXLWarning)]
    assert any("keep their formulas" in str(w.message) for w in ignored)


@pytest.mark.parametrize("engine", ENGINES)
def test_a_stale_definition_over_values_is_filled_as_data(
    stale_definition_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """A definition over constant values describes data, not a formula."""
    output = _fill_output(
        stale_definition_template,
        output_path,
        engine,
        tables={"Priced": [{"Name": "new", "Price": 42.5}]},
    )

    wb = openpyxl.load_workbook(output)
    try:
        assert wb["Data"]["B2"].value == 42.5
    finally:
        wb.close()


@pytest.mark.parametrize("engine", ENGINES)
def test_a_constant_calculated_column_is_kept(
    constant_formula_template: Path,
    output_path: Path,
    engine: Engine,
) -> None:
    """``=100`` is a valid calculated column and survives a data fill."""
    output = _fill_output(
        constant_formula_template,
        output_path,
        engine,
        tables={"Priced": [{"Name": "a"}, {"Name": "b"}]},
    )

    wb = openpyxl.load_workbook(output)
    try:
        ws = wb["Data"]
        for row in (2, 3):
            cell = ws.cell(row=row, column=2)
            assert cell.data_type == "f"
            assert cell.value == "=100"
    finally:
        wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# Totals rows
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.parametrize("engine", ENGINES)
def test_totals_row_shown_without_a_count_is_not_a_totals_row(
    tmp_path: Path,
    engine: Engine,
) -> None:
    """``totalsRowShown`` alone does not make the last row a totals row."""
    template = tmp_path / "totals_shown.xlsx"
    output = tmp_path / f"totals_shown_{engine}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Value"])
    ws.append(["old", 0])
    ws.append(["old", 1])
    ws.add_table(
        Table(
            displayName="Shown",
            ref="A1:B3",
            tableColumns=[TableColumn(id=1, name="Name"), TableColumn(id=2, name="Value")],
            totalsRowShown=True,
        ),
    )
    wb.save(template)
    wb.close()

    if engine == "win32com" and not _win32com_available():
        pytest.skip("win32com backend requires Windows and Excel")

    fill(
        template=template,
        output=output,
        engine=engine,
        tables={
            "Shown": [
                {"Name": "a", "Value": 1},
                {"Name": "b", "Value": 2},
                {"Name": "c", "Value": 3},
            ]
        },
    )

    result = openpyxl.load_workbook(output)
    try:
        sheet = result["Data"]
        table = sheet.tables["Shown"]
        assert table.ref == "A1:B4"
        assert [sheet.cell(row=row, column=1).value for row in (2, 3, 4)] == ["a", "b", "c"]
    finally:
        result.close()
