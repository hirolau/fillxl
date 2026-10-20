"""Cross-backend datatype consistency tests.

Verifies that fillxl handles common Python datatypes identically regardless
of whether the ``openpyxl`` or ``win32com`` backend is used.  Every dtype case
runs against both engines (``ENGINES``); the only Excel-only class holds
rendered-value checks that need xlwings.
"""

from __future__ import annotations

from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from typing import TYPE_CHECKING

import openpyxl
import pytest

from fillxl import Engine, FillXLConfig, fill
from tests._engines import ENGINES, _fill_output

if TYPE_CHECKING:
    from pathlib import Path

_PERMISSIVE = FillXLConfig(column_validation_strategy="none")

_NAMED_CELL_COL_MAP: dict[str, str] = {
    "cd_str_cell": "A",
    "cd_int_cell": "B",
    "cd_float_cell": "C",
    "cd_bool_cell": "D",
    "cd_none_cell": "E",
    "cd_datetime_cell": "F",
    "cd_date_cell": "G",
    "cd_time_cell": "H",
    "cd_decimal_cell": "I",
    "cd_nan_cell": "J",
    "cd_timestamp_cell": "K",
    "cd_tz_datetime_cell": "L",
    "cd_large_int_cell": "M",
    "cd_large_float_cell": "N",
}

_TABLE_COL_MAP: dict[str, str] = {
    "StrCol": "A",
    "IntCol": "B",
    "FloatCol": "C",
    "BoolCol": "D",
    "NoneCol": "E",
    "DatetimeCol": "F",
    "DateCol": "G",
    "TimeCol": "H",
    "DecimalCol": "I",
    "FloatNanCol": "J",
    "NaCol": "K",
    "NaTCol": "L",
    "TimestampCol": "M",
    "TzDatetimeCol": "N",
    "LargeIntCol": "O",
    "LargeFloatCol": "P",
}


def _read_value(path: Path, sheet: str, reference: str) -> object:
    """Read one cell back with openpyxl."""
    wb = openpyxl.load_workbook(path)
    try:
        return wb[sheet][reference].value
    finally:
        wb.close()


def _assert_calendar_day(value: object, expected: date) -> None:
    assert isinstance(value, datetime), f"expected datetime, got {type(value)}"
    assert (value.year, value.month, value.day) == (expected.year, expected.month, expected.day)


def _assert_time_of_day(value: object, expected: time) -> None:
    assert isinstance(value, time), f"expected time, got {type(value)}"
    assert (value.hour, value.minute, value.second) == (
        expected.hour,
        expected.minute,
        expected.second,
    )


# ═══════════════════════════════════════════════════════════════════════════════
# Named cells
# ═══════════════════════════════════════════════════════════════════════════════


class TestCrossBackendNamedCells:
    """Named cell dtype tests, run against both engines."""

    @pytest.mark.parametrize("engine", ENGINES)
    @pytest.mark.parametrize(
        ("cell_name", "value"),
        [
            pytest.param("cd_str_cell", "Hello, 世界", id="str"),
            pytest.param("cd_int_cell", 42, id="int"),
            pytest.param("cd_float_cell", 3.14159, id="float"),
            pytest.param("cd_bool_cell", True, id="bool_true"),
            pytest.param("cd_bool_cell", False, id="bool_false"),
            pytest.param("cd_none_cell", None, id="none"),
            pytest.param(
                "cd_datetime_cell",
                datetime(2025, 6, 15, 14, 30, 45),
                id="datetime",
            ),
            pytest.param("cd_date_cell", date(2025, 12, 25), id="date"),
            pytest.param("cd_time_cell", time(9, 15, 30), id="time"),
            pytest.param("cd_decimal_cell", Decimal("12.34"), id="decimal"),
            pytest.param("cd_nan_cell", float("nan"), id="float_nan"),
            pytest.param("cd_large_int_cell", 2**53 - 1, id="large_int"),
            pytest.param("cd_large_float_cell", 1e200, id="large_float"),
        ],
    )
    def test_named_cell_preserves_value(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
        cell_name: str,
        value: object,
    ) -> None:
        """A named cell round-trips the Python value on both engines."""
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            values={cell_name: value},
        )
        read_back = _read_value(
            output,
            "CrossDtypeNamedCells",
            f"{_NAMED_CELL_COL_MAP[cell_name]}1",
        )

        if cell_name in ("cd_none_cell", "cd_nan_cell"):
            assert read_back is None
        elif cell_name == "cd_datetime_cell":
            assert isinstance(value, datetime)
            _assert_calendar_day(read_back, value)
            assert isinstance(read_back, datetime)
            assert (read_back.hour, read_back.minute) == (value.hour, value.minute)
        elif cell_name == "cd_date_cell":
            assert isinstance(value, date)
            _assert_calendar_day(read_back, value)
        elif cell_name == "cd_time_cell":
            assert isinstance(value, time)
            _assert_time_of_day(read_back, value)
        elif cell_name == "cd_decimal_cell":
            assert read_back == pytest.approx(12.34)
        else:
            assert read_back == value

    @pytest.mark.parametrize("engine", ENGINES)
    def test_named_cell_converts_pandas_timestamp(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """A pd.Timestamp becomes a datetime on both engines."""
        pd = pytest.importorskip("pandas")
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            values={"cd_timestamp_cell": pd.Timestamp("2025-06-15 14:30:00")},
        )

        read_back = _read_value(output, "CrossDtypeNamedCells", "K1")
        _assert_calendar_day(read_back, date(2025, 6, 15))

    @pytest.mark.parametrize("engine", ENGINES)
    def test_named_cell_strips_timezone_from_datetime(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """A tz-aware datetime is stored naive and warns on both engines."""
        aware = datetime(2025, 6, 15, 14, 30, tzinfo=UTC)
        with pytest.warns(UserWarning, match="Timezone-aware"):
            output = _fill_output(
                cross_backend_dtype_template,
                output_path,
                engine,
                values={"cd_tz_datetime_cell": aware},
            )

        read_back = _read_value(output, "CrossDtypeNamedCells", "L1")
        _assert_calendar_day(read_back, date(2025, 6, 15))
        assert isinstance(read_back, datetime)
        assert (read_back.hour, read_back.minute) == (14, 30)
        assert read_back.tzinfo is None

    @pytest.mark.parametrize("engine", ENGINES)
    def test_named_cell_stores_timedelta_as_a_fraction_of_a_day(
        self,
        output_path: Path,
        engine: Engine,
    ) -> None:
        """``timedelta`` is normalised to a float on both engines."""
        from openpyxl.workbook.defined_name import DefinedName

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = ""
        wb.defined_names.add(DefinedName("dur", attr_text="Data!$A$1"))
        template = output_path.with_name("td_named_tmpl.xlsx")
        wb.save(template)
        wb.close()

        duration = timedelta(hours=12, minutes=30)
        output = _fill_output(
            template,
            output_path,
            engine,
            values={"dur": duration},
        )

        read_back = _read_value(output, "Data", "A1")
        assert read_back == pytest.approx(duration.total_seconds() / 86400.0)


# ═══════════════════════════════════════════════════════════════════════════════
# Tables
# ═══════════════════════════════════════════════════════════════════════════════


class TestCrossBackendTables:
    """Table column dtype tests, run against both engines."""

    @pytest.mark.parametrize("engine", ENGINES)
    @pytest.mark.parametrize(
        ("col_name", "values"),
        [
            pytest.param("StrCol", ["hello", "world"], id="str"),
            pytest.param("IntCol", [10, -5], id="int"),
            pytest.param("FloatCol", [1.5, -3.14], id="float"),
            pytest.param("BoolCol", [True, False], id="bool"),
            pytest.param("NoneCol", [None, None], id="none"),
            pytest.param(
                "DatetimeCol",
                [datetime(2025, 6, 15, 14, 30), datetime(2024, 1, 1, 0, 0)],
                id="datetime",
            ),
            pytest.param("DateCol", [date(2025, 12, 25), date(2023, 7, 4)], id="date"),
            pytest.param("TimeCol", [time(9, 15), time(23, 59, 59)], id="time"),
            pytest.param(
                "DecimalCol",
                [Decimal("12.34"), Decimal("0.01")],
                id="decimal",
            ),
            pytest.param("LargeIntCol", [2**53 - 1, -(2**53)], id="large_int"),
            pytest.param("LargeFloatCol", [1e200, 1e-300], id="large_float"),
        ],
    )
    def test_table_column_preserves_value(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
        col_name: str,
        values: list[object],
    ) -> None:
        """A table column round-trips its Python values on both engines."""
        data = [{col_name: value} for value in values]
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            tables={"CrossDtypeTable": data},
            config=_PERMISSIVE,
        )

        letter = _TABLE_COL_MAP[col_name]
        for index, value in enumerate(values):
            read_back = _read_value(output, "CrossDtypeTable", f"{letter}{index + 2}")
            _assert_table_value(col_name, read_back, value)

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_float_nan_is_written_as_none(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """NaN is an empty cell on both engines."""
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            tables={"CrossDtypeTable": [{"FloatNanCol": float("nan")}, {"FloatNanCol": 1.0}]},
            config=_PERMISSIVE,
        )

        assert _read_value(output, "CrossDtypeTable", "J2") is None
        assert _read_value(output, "CrossDtypeTable", "J3") == 1.0

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_pandas_na_is_written_as_none(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """pd.NA is an empty cell on both engines."""
        pd = pytest.importorskip("pandas")
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            tables={"CrossDtypeTable": [{"NaCol": pd.NA}, {"NaCol": "hello"}]},
            config=_PERMISSIVE,
        )

        assert _read_value(output, "CrossDtypeTable", "K2") is None
        assert _read_value(output, "CrossDtypeTable", "K3") == "hello"

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_pandas_nat_is_written_as_none(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """pd.NaT is an empty cell on both engines."""
        pd = pytest.importorskip("pandas")
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            tables={
                "CrossDtypeTable": [
                    {"NaTCol": pd.NaT},
                    {"NaTCol": datetime(2025, 1, 1)},
                ]
            },
            config=_PERMISSIVE,
        )

        assert _read_value(output, "CrossDtypeTable", "L2") is None
        assert _read_value(output, "CrossDtypeTable", "L3") is not None

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_pandas_timestamp_becomes_datetime(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """pd.Timestamp becomes a datetime on both engines."""
        pd = pytest.importorskip("pandas")
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            tables={
                "CrossDtypeTable": [
                    {"TimestampCol": pd.Timestamp("2025-06-15 14:30:00")},
                    {"TimestampCol": pd.Timestamp("2024-01-01")},
                ]
            },
            config=_PERMISSIVE,
        )

        _assert_calendar_day(_read_value(output, "CrossDtypeTable", "M2"), date(2025, 6, 15))
        _assert_calendar_day(_read_value(output, "CrossDtypeTable", "M3"), date(2024, 1, 1))

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_tz_aware_timestamp_is_stripped_and_warns(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """A tz-aware pandas timestamp is stored naive and warns on both engines."""
        pd = pytest.importorskip("pandas")
        with pytest.warns(UserWarning, match="Timezone-aware"):
            output = _fill_output(
                cross_backend_dtype_template,
                output_path,
                engine,
                tables={
                    "CrossDtypeTable": [
                        {"TzDatetimeCol": pd.Timestamp("2025-06-15 14:30:00", tz="UTC")}
                    ]
                },
                config=_PERMISSIVE,
            )

        read_back = _read_value(output, "CrossDtypeTable", "N2")
        _assert_calendar_day(read_back, date(2025, 6, 15))
        assert isinstance(read_back, datetime)
        assert (read_back.hour, read_back.minute) == (14, 30)
        assert read_back.tzinfo is None

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_writes_mixed_dtypes_in_single_row(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        engine: Engine,
    ) -> None:
        """One row containing every dtype is written successfully on both engines."""
        data = [
            {
                "StrCol": "text",
                "IntCol": 99,
                "FloatCol": 3.14,
                "BoolCol": True,
                "NoneCol": None,
                "DatetimeCol": datetime(2025, 6, 15, 12, 0),
                "DateCol": date(2025, 6, 15),
                "TimeCol": time(12, 0),
                "DecimalCol": Decimal("99.99"),
                "FloatNanCol": float("nan"),
                "LargeIntCol": 2**53 - 1,
                "LargeFloatCol": 1e200,
            },
        ]
        output = _fill_output(
            cross_backend_dtype_template,
            output_path,
            engine,
            tables={"CrossDtypeTable": data},
            config=_PERMISSIVE,
        )

        assert _read_value(output, "CrossDtypeTable", "A2") == "text"
        assert _read_value(output, "CrossDtypeTable", "B2") == 99
        assert _read_value(output, "CrossDtypeTable", "C2") == 3.14
        assert bool(_read_value(output, "CrossDtypeTable", "D2")) is True
        assert _read_value(output, "CrossDtypeTable", "E2") is None
        _assert_calendar_day(_read_value(output, "CrossDtypeTable", "F2"), date(2025, 6, 15))
        _assert_calendar_day(_read_value(output, "CrossDtypeTable", "G2"), date(2025, 6, 15))
        _assert_time_of_day(_read_value(output, "CrossDtypeTable", "H2"), time(12, 0))
        assert _read_value(output, "CrossDtypeTable", "I2") == pytest.approx(99.99)
        assert _read_value(output, "CrossDtypeTable", "J2") is None
        assert _read_value(output, "CrossDtypeTable", "O2") == 2**53 - 1
        assert _read_value(output, "CrossDtypeTable", "P2") == 1e200

    @pytest.mark.parametrize("engine", ENGINES)
    def test_table_stores_timedelta_as_a_fraction_of_a_day(
        self,
        output_path: Path,
        engine: Engine,
    ) -> None:
        """``timedelta`` in a table cell is a float on both engines."""
        from openpyxl.worksheet.table import Table, TableColumn

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "Duration"
        ws["A2"] = None
        ws.add_table(
            Table(
                displayName="Times",
                ref="A1:A2",
                tableColumns=[TableColumn(id=1, name="Duration")],
            )
        )
        template = output_path.with_name("td_table_tmpl.xlsx")
        wb.save(template)
        wb.close()

        duration = timedelta(hours=1, minutes=30)
        output = _fill_output(
            template,
            output_path,
            engine,
            tables={"Times": [{"Duration": duration}]},
            config=_PERMISSIVE,
        )

        read_back = _read_value(output, "Data", "A2")
        assert read_back == pytest.approx(duration.total_seconds() / 86400.0)


def _assert_table_value(col_name: str, read_back: object, value: object) -> None:
    """Assert one table cell's read-back value for its dtype."""
    if col_name == "NoneCol":
        assert read_back is None
    elif col_name == "BoolCol":
        assert bool(read_back) is value
    elif col_name == "DatetimeCol":
        assert isinstance(value, datetime)
        _assert_calendar_day(read_back, value)
    elif col_name == "DateCol":
        assert isinstance(value, date)
        _assert_calendar_day(read_back, value)
    elif col_name == "TimeCol":
        assert isinstance(value, time)
        _assert_time_of_day(read_back, value)
    elif col_name == "DecimalCol":
        assert isinstance(value, Decimal)
        assert read_back == pytest.approx(float(value))
    else:
        assert read_back == value


# ═══════════════════════════════════════════════════════════════════════════════
# Excel-only: rendered values through xlwings
# ═══════════════════════════════════════════════════════════════════════════════


@pytest.mark.win32com
class TestCrossBackendXlwingsReadBack:
    """Rendered value verification via xlwings COM read-back (Excel required).

    These tests open the output in real Excel via xlwings to verify the
    *rendered* values; no other test can observe Excel's display layer.
    """

    def test_named_cell_values_match_via_xlwings(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        xl_app: object,  # xw.App
    ) -> None:
        """Named cell values rendered by Excel match expected values (xlwings read-back)."""
        import xlwings as xw

        fill(
            template=cross_backend_dtype_template,
            output=output_path,
            values={
                "cd_str_cell": "Excel",
                "cd_int_cell": 99,
                "cd_float_cell": 2.71828,
                "cd_bool_cell": True,
                "cd_none_cell": None,
                "cd_datetime_cell": datetime(2025, 6, 15, 14, 30, 0),
            },
            engine="win32com",
        )

        wb = xw.Book(str(output_path))
        try:
            ws = wb.sheets["CrossDtypeNamedCells"]
            assert ws.range("A1").value == "Excel"
            assert ws.range("B1").value == 99
            assert abs(ws.range("C1").value - 2.71828) < 1e-5
            assert ws.range("D1").value in (True, 1.0, 1)
            assert ws.range("E1").value is None
            assert ws.range("F1").value is not None
        finally:
            wb.close()

    def test_table_values_match_via_xlwings(
        self,
        output_path: Path,
        cross_backend_dtype_template: Path,
        xl_app: object,
    ) -> None:
        """Table values rendered by Excel match expected values (xlwings read-back)."""
        import xlwings as xw

        fill(
            template=cross_backend_dtype_template,
            output=output_path,
            tables={
                "CrossDtypeTable": [
                    {
                        "StrCol": "hello",
                        "IntCol": 42,
                        "FloatCol": 3.14,
                        "BoolCol": True,
                        "NoneCol": None,
                    }
                ]
            },
            config=_PERMISSIVE,
            engine="win32com",
        )

        wb = xw.Book(str(output_path))
        try:
            ws = wb.sheets["CrossDtypeTable"]
            assert ws.range("A2").value == "hello"
            assert ws.range("B2").value == 42
            assert abs(ws.range("C2").value - 3.14) < 1e-5
        finally:
            wb.close()
