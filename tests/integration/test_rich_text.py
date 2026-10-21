"""Partly formatted text in a template stays formatted."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from fillxl import fill
from tests.parity.builders import table_template

_ROWS = [{"fund": f"F{i}", "nav": i} for i in range(4)]


def _rich_text() -> CellRichText:
    return CellRichText([TextBlock(InlineFont(b=True), "Bold"), " plain"])


def test_partly_bold_text_in_the_template_survives_the_fill(tmp_path: Path) -> None:
    """openpyxl flattens rich text to a plain string unless asked to keep it."""
    template = table_template(tmp_path / "template.xlsx", top_row=3)
    workbook = openpyxl.load_workbook(template, rich_text=True)
    workbook["Data"]["A1"] = _rich_text()
    workbook.save(template)
    workbook.close()

    output = tmp_path / "out.xlsx"
    fill(template=template, output=output, engine="openpyxl", tables={"Sales": _ROWS})

    workbook = openpyxl.load_workbook(output, rich_text=True)
    try:
        value = workbook["Data"]["A1"].value
        assert isinstance(value, CellRichText)
        assert value[0].font.b  # type: ignore[union-attr]
        assert str(value) == "Bold plain"
    finally:
        workbook.close()


def test_rich_text_in_a_totals_row_survives_a_resize(tmp_path: Path) -> None:
    """The totals row is recreated below the new body, formatting included."""
    template = table_template(tmp_path / "template.xlsx", body=(("A", 1), ("B", 2)), totals=True)
    workbook = openpyxl.load_workbook(template, rich_text=True)
    worksheet = workbook["Data"]
    worksheet["A4"] = _rich_text()
    workbook.save(template)
    workbook.close()

    output = tmp_path / "out.xlsx"
    fill(template=template, output=output, engine="openpyxl", tables={"Sales": _ROWS})

    workbook = openpyxl.load_workbook(output, rich_text=True)
    try:
        worksheet = workbook["Data"]
        table = worksheet.tables["Sales"]
        assert table.ref == "A1:B6"
        value = worksheet["A6"].value
        assert isinstance(value, CellRichText)
        assert value[0].font.b  # type: ignore[union-attr]
    finally:
        workbook.close()


@pytest.mark.parametrize("column", [1, 2])
def test_a_cell_outside_the_table_is_never_touched(tmp_path: Path, column: int) -> None:
    """Only the table body is rewritten."""
    template = table_template(tmp_path / "template.xlsx")
    workbook = openpyxl.load_workbook(template, rich_text=True)
    workbook["Data"].cell(row=10, column=column, value=_rich_text())
    workbook.save(template)
    workbook.close()

    output = tmp_path / "out.xlsx"
    fill(template=template, output=output, engine="openpyxl", tables={"Sales": _ROWS})

    workbook = openpyxl.load_workbook(output, rich_text=True)
    try:
        assert isinstance(workbook["Data"].cell(row=10, column=column).value, CellRichText)
    finally:
        workbook.close()
