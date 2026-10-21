"""Sheet visibility, deletion and the active sheet, in the openpyxl engine."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pytest
from openpyxl.workbook.protection import WorkbookProtection

from fillxl import FillXLConfig, fill
from tests.parity.builders import chart_sheet_first_template, sheets_template


def _active(path: Path) -> str | None:
    workbook = openpyxl.load_workbook(path)
    try:
        active = workbook.active
        return None if active is None else str(active.title)
    finally:
        workbook.close()


def _fill(template: Path, output: Path, **config: Any) -> Path:  # noqa: ANN401
    return fill(
        template=template,
        output=output,
        engine="openpyxl",
        config=FillXLConfig(**config),
    )


def test_hiding_the_active_underscore_sheet_makes_the_first_visible_sheet_active(
    tmp_path: Path,
) -> None:
    """A hidden sheet must not stay the active one."""
    template = sheets_template(tmp_path / "template.xlsx", "_data", "Report")

    output = _fill(template, tmp_path / "out.xlsx")

    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook["_data"].sheet_state == "hidden"
        assert workbook.active is not None
        assert workbook.active.title == "Report"
    finally:
        workbook.close()


def test_a_very_hidden_sheet_stays_very_hidden(tmp_path: Path) -> None:
    """``veryHidden`` sheets are not in Excel's Unhide dialog, so they stay that way."""
    template = sheets_template(tmp_path / "template.xlsx", "_data", "Report")
    workbook = openpyxl.load_workbook(template)
    workbook["_data"].sheet_state = "veryHidden"
    workbook.save(template)
    workbook.close()

    output = _fill(template, tmp_path / "out.xlsx")

    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook["_data"].sheet_state == "veryHidden"
    finally:
        workbook.close()


def test_hiding_every_sheet_is_refused_instead_of_writing_a_broken_file(
    tmp_path: Path,
) -> None:
    """Excel requires one visible sheet, so the fill fails instead of corrupting the file."""
    template = sheets_template(tmp_path / "template.xlsx", "_data")

    with pytest.raises(ValueError, match="without a visible sheet"):
        _fill(template, tmp_path / "out.xlsx")

    assert not (tmp_path / "out.xlsx").exists()


def test_deleting_the_active_sheet_makes_the_first_remaining_visible_sheet_active(
    tmp_path: Path,
) -> None:
    """openpyxl keeps the active *index*, which now points at another sheet."""
    template = sheets_template(tmp_path / "template.xlsx", "Report", "Summary", active="Report")

    output = _fill(
        template, tmp_path / "out.xlsx", hide_underscore_sheets=False, delete_sheets=["Report"]
    )

    assert _active(output) == "Summary"
    assert openpyxl.load_workbook(output).sheetnames == ["Summary"]


def test_deleting_a_sheet_before_the_active_one_keeps_the_same_sheet_active(
    tmp_path: Path,
) -> None:
    """The active sheet does not change just because something before it went."""
    template = sheets_template(tmp_path / "template.xlsx", "Report", "Summary", active="Summary")

    output = _fill(
        template, tmp_path / "out.xlsx", hide_underscore_sheets=False, delete_sheets=["Report"]
    )

    assert _active(output) == "Summary"


def test_deleting_every_visible_sheet_is_refused(tmp_path: Path) -> None:
    """A hidden sheet does not count as the one Excel needs."""
    template = sheets_template(tmp_path / "template.xlsx", "_data", "Report")
    workbook = openpyxl.load_workbook(template)
    workbook["_data"].sheet_state = "hidden"
    workbook.save(template)
    workbook.close()

    with pytest.raises(ValueError, match="without a visible sheet"):
        _fill(template, tmp_path / "out.xlsx", delete_sheets=["Report"])


def test_a_hidden_sheet_cannot_be_made_active(tmp_path: Path) -> None:
    """Asking for a hidden active sheet is a mistake worth reporting."""
    template = sheets_template(tmp_path / "template.xlsx", "_data", "Report")

    with pytest.raises(ValueError, match="hidden and cannot be the active sheet"):
        _fill(template, tmp_path / "out.xlsx", active_sheet="_data")


def test_choosing_the_active_sheet_leaves_exactly_one_tab_selected(tmp_path: Path) -> None:
    """Two selected tabs make Excel open the workbook with sheets grouped."""
    template = sheets_template(tmp_path / "template.xlsx", "Report", "Summary")

    output = _fill(template, tmp_path / "out.xlsx", active_sheet="Summary")

    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook.active is not None
        assert workbook.active.title == "Summary"
        selected = [name for name in workbook.sheetnames if workbook[name].sheet_view.tabSelected]
        assert selected == ["Summary"]
    finally:
        workbook.close()


def test_the_active_sheet_is_found_by_name_even_after_a_chart_sheet(tmp_path: Path) -> None:
    """A worksheet index is not a position in ``sheetnames`` when chart sheets come first."""
    template = chart_sheet_first_template(tmp_path / "template.xlsx")

    output = _fill(template, tmp_path / "out.xlsx", active_sheet="Report")

    assert _active(output) == "Report"


def test_a_chart_sheet_can_be_deleted(tmp_path: Path) -> None:
    """Chart sheets used to be skipped, which left them in the report."""
    template = chart_sheet_first_template(tmp_path / "template.xlsx")

    output = _fill(
        template, tmp_path / "out.xlsx", hide_underscore_sheets=False, delete_sheets=["Chart"]
    )

    assert openpyxl.load_workbook(output).sheetnames == ["Report"]


@pytest.mark.excel
def test_a_failed_delete_is_reported_not_mistaken_for_a_missing_sheet(
    tmp_path: Path,
) -> None:
    """A protected structure makes Excel refuse; that is not "not found"."""
    template = sheets_template(tmp_path / "template.xlsx", "Data", "Report")
    workbook = openpyxl.load_workbook(template)
    workbook.security = WorkbookProtection(workbookPassword="x", lockStructure=True)
    workbook.save(template)
    workbook.close()

    with pytest.raises(RuntimeError, match="Could not delete sheet 'Report'"):
        fill(
            template=template,
            output=tmp_path / "out.xlsx",
            engine="win32com",
            config=FillXLConfig(hide_underscore_sheets=False, delete_sheets=["Report"]),
        )
