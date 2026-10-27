"""What happens around a resized table: its neighbours, merged cells, sort state.

A table that grows must not silently overwrite the workbook around it, and
everything that pointed at the old body must follow the new one.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableColumn

from fillxl import fill
from tests._engines import ENGINES, _fill_output
from tests.parity.builders import table_template

if TYPE_CHECKING:
    from pathlib import Path

    from fillxl import Engine

# ═══════════════════════════════════════════════════════════════════════════════
# Growing into something
# ═══════════════════════════════════════════════════════════════════════════════


def _table(ref: str, name: str, columns: tuple[str, ...]) -> Table:
    return Table(
        displayName=name,
        ref=ref,
        tableColumns=[TableColumn(id=i, name=c) for i, c in enumerate(columns, start=1)],
    )


def _two_tables_template(path: Path) -> Path:
    """``Sales`` at A1:B2 with ``Other`` sitting directly below it at A4:B5."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet["A1"], worksheet["B1"] = "fund", "nav"
    worksheet["A2"], worksheet["B2"] = "A", 1
    # The header cells must spell the column names, or Excel refuses the file.
    worksheet["A4"], worksheet["B4"] = "name", "value"
    worksheet["A5"], worksheet["B5"] = "other", 3
    worksheet.add_table(_table("A1:B2", "Sales", ("fund", "nav")))
    worksheet.add_table(_table("A4:B5", "Other", ("name", "value")))
    workbook.save(path)
    return path


@pytest.mark.parametrize("engine", ENGINES)
def test_a_table_refuses_to_grow_into_another_table(
    tmp_path: Path,
    engine: Engine,
) -> None:
    """An overlap makes Excel repair the workbook, so it is refused up front."""
    template = _two_tables_template(tmp_path / "template.xlsx")
    rows = [{"fund": letter, "nav": index} for index, letter in enumerate("ABCD", start=1)]

    with pytest.raises(ValueError, match=r"table 'Other' at A4:B5 is in the way"):
        _fill_output(template, tmp_path / "out.xlsx", engine, tables={"Sales": rows})


def _merged_below_template(path: Path) -> Path:
    """``Sales`` at A1:B2 with an empty merged block at A5:B5."""
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet["A1"], worksheet["B1"] = "fund", "nav"
    worksheet["A2"], worksheet["B2"] = "A", 1
    worksheet.merge_cells("A5:B5")
    worksheet.add_table(_table("A1:B2", "Sales", ("fund", "nav")))
    workbook.save(path)
    return path


@pytest.mark.parametrize("engine", ENGINES)
def test_a_table_refuses_to_grow_into_merged_cells(
    tmp_path: Path,
    engine: Engine,
) -> None:
    """The cells look empty but writing to one raises openpyxl's read-only error."""
    template = _merged_below_template(tmp_path / "template.xlsx")
    rows = [{"fund": letter, "nav": index} for index, letter in enumerate("ABCD", start=1)]

    with pytest.raises(ValueError, match=r"the merged range A5:B5 is in the way"):
        _fill_output(template, tmp_path / "out.xlsx", engine, tables={"Sales": rows})


# ═══════════════════════════════════════════════════════════════════════════════
# Sort state
# ═══════════════════════════════════════════════════════════════════════════════


def _sorted_template(path: Path) -> Path:
    """A table saved while sorted: the sort state still covers the old body."""
    from openpyxl.worksheet.filters import SortCondition, SortState

    template = table_template(
        path,
        body=tuple(("fund", index) for index in (1, 2, 3)),
        name="Sales",
    )
    workbook = openpyxl.load_workbook(template)
    table = workbook["Data"].tables["Sales"]
    table.sortState = SortState(
        ref="A2:B4", sortCondition=[SortCondition(ref="A2:A4", sortBy="value")]
    )
    workbook.save(template)
    workbook.close()
    return template


@pytest.mark.parametrize("engine", ENGINES)
def test_a_sorted_table_keeps_the_sort_state_the_template_had(
    tmp_path: Path,
    engine: Engine,
) -> None:
    """Neither engine rewrites ``sortState``, because Excel does not either.

    Measured in Excel 16: resizing a table (in either direction) leaves the
    old ``sortState`` ref in place, and the workbook still opens without a
    repair prompt.  Rewriting it here would make the two engines disagree.
    """
    template = _sorted_template(tmp_path / "template.xlsx")

    output = _fill_output(
        template,
        tmp_path / "out.xlsx",
        engine,
        tables={"Sales": [{"fund": f"f{i}", "nav": i} for i in range(6)]},
    )

    workbook = openpyxl.load_workbook(output)
    try:
        sort_state = workbook["Data"].tables["Sales"].sortState
        assert sort_state is not None
        assert sort_state.ref == "A2:B4"
        assert [condition.ref for condition in sort_state.sortCondition] == ["A2:A4"]
    finally:
        workbook.close()


# ═══════════════════════════════════════════════════════════════════════════════
# Calculated columns and totals rows
# ═══════════════════════════════════════════════════════════════════════════════


def _calculated_template(path: Path) -> Path:
    """A table whose third column is a calculated column."""
    from openpyxl.worksheet.table import TableFormula

    template = table_template(
        path,
        header=("fund", "nav", "double"),
        body=tuple((f"f{index}", index, None) for index in range(1, 5)),
        name="Sales",
    )
    workbook = openpyxl.load_workbook(template)
    worksheet = workbook["Data"]
    table = worksheet.tables["Sales"]
    table.tableColumns[2].calculatedColumnFormula = TableFormula(
        attr_text="Sales[[#This Row],[nav]]*2"
    )
    for row in range(2, 6):
        worksheet.cell(row=row, column=3, value="=Sales[[#This Row],[nav]]*2")
    workbook.save(template)
    workbook.close()
    return template


@pytest.mark.parametrize("engine", ENGINES)
def test_empty_data_keeps_the_calculated_column_formula(
    tmp_path: Path,
    engine: Engine,
) -> None:
    """Excel keeps the formula in the single blank row, so fillxl keeps it too."""
    import pandas as pd

    template = _calculated_template(tmp_path / "template.xlsx")

    output = _fill_output(
        template,
        tmp_path / "out.xlsx",
        engine,
        tables={"Sales": pd.DataFrame(columns=["fund", "nav"])},
    )

    workbook = openpyxl.load_workbook(output)
    try:
        worksheet = workbook["Data"]
        table = worksheet.tables["Sales"]
        assert table.ref == "A1:C2"
        assert worksheet["C2"].value is not None
        assert str(worksheet["C2"].value).startswith("=")
        definition = table.tableColumns[2].calculatedColumnFormula
        assert definition is not None
    finally:
        workbook.close()


def _totals_template(path: Path) -> Path:
    """A table with five body rows and a bold totals row."""
    template = table_template(
        path,
        header=("fund", "nav"),
        body=tuple((f"f{index}", index) for index in range(1, 6)),
        name="Sales",
        totals=True,
    )
    workbook = openpyxl.load_workbook(template)
    worksheet = workbook["Data"]
    for column in ("A", "B"):
        worksheet[f"{column}7"].font = Font(bold=True)
    workbook.save(template)
    workbook.close()
    return template


def test_shrinking_a_totals_table_takes_the_old_totals_formatting_away(
    tmp_path: Path,
) -> None:
    """The totals row moves up; its bold text must not stay behind."""
    template = _totals_template(tmp_path / "template.xlsx")

    output = _fill_output(
        template,
        tmp_path / "out.xlsx",
        "openpyxl",
        tables={"Sales": [{"fund": "only", "nav": 1}]},
    )

    workbook = openpyxl.load_workbook(output)
    try:
        worksheet = workbook["Data"]
        assert worksheet["A3"].font.bold, "the totals row should have moved to row 3"
        assert not worksheet["A7"].font.bold, "the old totals row kept its formatting"
    finally:
        workbook.close()


# ═══════════════════════════════════════════════════════════════════════════════
# Plain fill() guards
# ═══════════════════════════════════════════════════════════════════════════════


def test_the_template_is_never_overwritten(tmp_path: Path) -> None:
    """A report pipeline that points output at its own template is a mistake."""
    template = table_template(tmp_path / "template.xlsx")

    with pytest.raises(ValueError, match="never overwrites its template"):
        fill(template, template, tables={"Sales": [{"fund": "A", "nav": 1}]})

    assert template.exists()
