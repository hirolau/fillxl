"""Workbook builders for the parity tests."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo


def table_template(
    path: Path,
    *,
    header: tuple[str, ...] = ("fund", "nav"),
    header_cells: tuple[Any, ...] | None = None,
    body: tuple[tuple[Any, ...], ...] = (("A", 1),),
    sheet: str = "Data",
    name: str = "Sales",
    top_row: int = 1,
    totals: bool = False,
    headerless: bool = False,
    number_formats: dict[int, str] | None = None,
    other_sheets: tuple[str, ...] = (),
    active_sheet: str | None = None,
) -> Path:
    """A workbook with one table. *header* holds the table column names as stored in the file;
    *header_cells* (default: *header*) holds what the header cells display."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for other in other_sheets:
        workbook.create_sheet(other)

    row = top_row
    if not headerless:
        for col, title in enumerate(header_cells or header, start=1):
            worksheet.cell(row=row, column=col, value=title)
        row += 1
    for values in body:
        for col, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row, column=col, value=value)
            if number_formats and col in number_formats:
                cell.number_format = number_formats[col]
        row += 1
    body_last_row = row - 1

    columns = [TableColumn(id=i, name=title) for i, title in enumerate(header, start=1)]
    last_row = body_last_row
    last_col = get_column_letter(len(header))
    if totals:
        last_row += 1
        worksheet.cell(row=last_row, column=1, value="Total")
        worksheet.cell(
            row=last_row, column=len(header), value=f"=SUBTOTAL(109,{name}[{header[-1]}])"
        )
        columns[0].totalsRowLabel = "Total"
        columns[-1].totalsRowFunction = "sum"

    table = Table(
        displayName=name,
        ref=f"A{top_row}:{last_col}{last_row}",
        tableColumns=columns,
        headerRowCount=0 if headerless else 1,
        totalsRowCount=1 if totals else None,
        tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
    )
    if not headerless:
        table.autoFilter = AutoFilter(ref=f"A{top_row}:{last_col}{body_last_row}")
    worksheet.add_table(table)

    if active_sheet is not None:
        index = workbook.sheetnames.index(active_sheet)
        workbook.active = index
        for position, sheet_name in enumerate(workbook.sheetnames):
            workbook[sheet_name].sheet_view.tabSelected = position == index
    workbook.save(path)
    return path


def named_cells_template(
    path: Path,
    *,
    cells: dict[str, tuple[str, Any]],
    sheet: str = "Data",
) -> Path:
    """A workbook whose *cells* map a defined name to ``(coordinate, value)``."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    for name, (coordinate, value) in cells.items():
        worksheet[coordinate] = value
        workbook.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!{coordinate}"))
    workbook.save(path)
    return path
