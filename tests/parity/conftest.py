"""Fill a template with both engines and compare the results cell by cell."""

from __future__ import annotations

import contextlib
import functools
import importlib
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText
from openpyxl.utils.escape import unescape
from openpyxl.worksheet.worksheet import Worksheet

import fillxl
from fillxl._charts import REFERENCE_RE, _unquote_sheet_name

_MISSING = object()
_SERIES_PARTS = ("cat", "val", "xVal", "yVal")
_FORMAT_ESCAPE_RE = re.compile(r"\\(.)")


def _number_format(cell: Any) -> str:  # noqa: ANN401
    """The cell's number format without Excel's backslash escapes.

    Excel escapes literal characters, so it writes the ISO date format as
    ``yyyy\\-mm\\-dd`` where openpyxl writes ``yyyy-mm-dd``.  The two spell
    the same format.
    """
    return _FORMAT_ESCAPE_RE.sub(r"\1", str(cell.number_format))


@functools.cache
def excel_is_available() -> bool:
    """Whether Microsoft Excel can be started through COM on this machine.

    The COM modules are imported by name: they only exist on Windows, and a
    plain ``import`` inside the guard leaves a name a static analyser cannot
    prove is bound on the platform where they are missing.
    """
    if sys.platform != "win32":
        return False
    try:
        pythoncom = importlib.import_module("pythoncom")
        win32com_client = importlib.import_module("win32com.client")
        pythoncom.CoInitialize()
        app = win32com_client.DispatchEx("Excel.Application")
        # A quit failure must not turn a usable Excel into "no Excel".
        with contextlib.suppress(Exception):
            app.Quit()
    except Exception:  # any failure means "no usable Excel"
        return False
    return True


def pytest_collection_modifyitems(
    config: pytest.Config,  # noqa: ARG001 - the hook signature is fixed
    items: list[pytest.Item],
) -> None:
    if excel_is_available():
        return
    skip = pytest.mark.skip(reason="needs Windows with Microsoft Excel")
    for item in items:
        if "excel" in item.keywords:
            item.add_marker(skip)


def _normalize_references(text: str) -> str:
    """Write every sheet reference unquoted so Excel's quoting choices don't matter."""
    return REFERENCE_RE.sub(
        lambda match: f"{_unquote_sheet_name(match['sheet'])}!{match['area']}", text
    )


def _cell(cell: Any) -> tuple[Any, ...]:  # noqa: ANN401
    value = cell.value
    if cell.data_type == "f":
        text = str(getattr(value, "text", value))  # ArrayFormula keeps its text in .text
        value = ("formula", text.replace("_xlfn.", "").replace("_xlpm.", "").replace(" ", ""))
    elif isinstance(value, CellRichText):
        value = ("rich text", str(value))
    return (value, cell.data_type, _number_format(cell))


def _chart_references(sheet: Any) -> dict[str, list[str]]:  # noqa: ANN401
    charts: dict[str, list[str]] = {}
    for chart_index, chart in enumerate(getattr(sheet, "_charts", [])):
        references: list[str] = []
        for series in chart.series:
            for part in _SERIES_PARTS:
                source = getattr(series, part, None)
                for kind in ("numRef", "strRef"):
                    reference = getattr(source, kind, None) if source is not None else None
                    if reference is not None and reference.f:
                        references.append(f"{part}={_normalize_references(reference.f)}")
        charts[f"chart{chart_index}"] = references
    return charts


def _sheet(sheet: Any) -> dict[str, Any]:  # noqa: ANN401
    snapshot: dict[str, Any] = {"state": sheet.sheet_state, "charts": _chart_references(sheet)}
    if not isinstance(sheet, Worksheet):  # a chart sheet
        return snapshot
    snapshot["cells"] = {
        cell.coordinate: _cell(cell)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    snapshot["tables"] = {
        table.name: {
            "ref": table.ref,
            "autofilter": table.autoFilter.ref if table.autoFilter is not None else None,
            "sortstate": table.sortState.ref if table.sortState is not None else None,
            "totals": table.totalsRowCount or 0,
            "columns": [unescape(str(column.name)) for column in table.tableColumns],
        }
        for table in sheet.tables.values()
    }
    snapshot["conditional_formatting"] = sorted(
        " ".join(sorted(str(r) for r in cf.sqref.ranges)) for cf in sheet.conditional_formatting
    )
    snapshot["names"] = {
        name: _normalize_references(str(defined.attr_text))
        for name, defined in sheet.defined_names.items()
    }
    return snapshot


def snapshot(path: Path) -> dict[str, Any]:
    """Everything a reader of the filled workbook would notice, as plain data."""
    workbook = openpyxl.load_workbook(path, rich_text=True)
    try:
        active = workbook.active
        return {
            "active": active.title if active is not None else None,
            "order": list(workbook.sheetnames),
            "sheets": {name: _sheet(workbook[name]) for name in workbook.sheetnames},
            "names": {
                name: _normalize_references(str(defined.attr_text))
                for name, defined in workbook.defined_names.items()
                # Excel adds this hidden name for sheet autofilters; it is not template content.
                if not name.startswith("_xlnm._FilterDatabase")
            },
        }
    finally:
        workbook.close()


def _flatten(value: Any, path: str = "") -> list[tuple[str, Any]]:  # noqa: ANN401
    if isinstance(value, dict):
        items: list[tuple[str, Any]] = []
        for key, item in value.items():
            items.extend(_flatten(item, f"{path}/{key}"))
        return items
    return [(path, value)]


class Parity:
    """Fills a template with both engines and compares the outputs."""

    def __init__(self, tmp_path: Path) -> None:
        self._tmp_path = tmp_path

    def fill_both(self, template: Path, **kwargs: Any) -> tuple[Path, Path]:  # noqa: ANN401
        outputs = [
            fillxl.fill(
                template,
                self._tmp_path / engine / f"filled{template.suffix}",
                engine=engine,
                **kwargs,
            )
            for engine in ("openpyxl", "win32com")
        ]
        return outputs[0], outputs[1]

    def assert_same(
        self, openpyxl_output: Path, excel_output: Path, *, ignore: tuple[str, ...] = ()
    ) -> None:
        """Every entry in *ignore* is a path prefix and MUST have a comment saying why."""
        left = dict(_flatten(snapshot(openpyxl_output)))
        right = dict(_flatten(snapshot(excel_output)))
        differences = [
            f"{key}: openpyxl={left.get(key, '<missing>')!r} excel={right.get(key, '<missing>')!r}"
            for key in sorted(left.keys() | right.keys())
            if not key.startswith(ignore) and left.get(key, _MISSING) != right.get(key, _MISSING)
        ]
        assert not differences, "The engines disagree:\n" + "\n".join(differences[:50])


@pytest.fixture
def parity(tmp_path: Path) -> Parity:
    return Parity(tmp_path)
