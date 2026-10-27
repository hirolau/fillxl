"""The same template and data must give the same workbook in both engines."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest

from fillxl import FillXLConfig, fill
from tests.parity.builders import named_cells_template, table_template
from tests.parity.conftest import Parity

pytestmark = pytest.mark.excel


def test_a_growing_table_ends_up_the_same_whichever_engine_fills_it(
    tmp_path: Path, parity: Parity
) -> None:
    template = table_template(tmp_path / "template.xlsx")
    rows = [{"fund": f"F{i}", "nav": i * 1.5} for i in range(10)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))


def test_a_shrinking_table_leaves_nothing_behind_in_either_engine(
    tmp_path: Path, parity: Parity
) -> None:
    body = tuple((f"F{i}", i) for i in range(8))
    template = table_template(tmp_path / "template.xlsx", body=body)
    parity.assert_same(*parity.fill_both(template, tables={"Sales": [{"fund": "A", "nav": 1}]}))


def test_an_empty_payload_clears_the_table_the_same_way_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    body = tuple((f"F{i}", i) for i in range(4))
    template = table_template(tmp_path / "template.xlsx", body=body)
    parity.assert_same(*parity.fill_both(template, tables={"Sales": []}))


def test_a_table_with_a_totals_row_grows_the_same_way_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    template = table_template(tmp_path / "template.xlsx", totals=True)
    rows = [{"fund": f"F{i}", "nav": i + 1} for i in range(6)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))


def test_a_table_with_a_totals_row_shrinks_the_same_way_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    body = tuple((f"F{i}", i + 1) for i in range(6))
    template = table_template(tmp_path / "template.xlsx", body=body, totals=True)
    parity.assert_same(*parity.fill_both(template, tables={"Sales": [{"fund": "A", "nav": 3}]}))


def test_a_headerless_table_grows_the_same_way_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    template = table_template(tmp_path / "template.xlsx", headerless=True)
    rows = [{"fund": f"F{i}", "nav": i * 2} for i in range(5)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))


def test_a_headerless_table_shrinks_the_same_way_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    body = tuple((f"F{i}", i) for i in range(6))
    template = table_template(tmp_path / "template.xlsx", body=body, headerless=True)
    parity.assert_same(*parity.fill_both(template, tables={"Sales": [{"fund": "A", "nav": 9}]}))


@pytest.mark.parametrize("escape_formulas", [True, False])
def test_strings_excel_likes_to_reinterpret_are_stored_the_same_way(
    tmp_path: Path, parity: Parity, escape_formulas: bool
) -> None:
    template = table_template(tmp_path / "template.xlsx", header=("text",), body=(("x",),))
    values = ["00123", "1/2", "#N/A", "TRUE", "=1+1"]
    rows = [{"text": value} for value in values]
    parity.assert_same(
        *parity.fill_both(
            template,
            tables={"Sales": rows},
            config=FillXLConfig(escape_formulas=escape_formulas),
        )
    )


def test_named_cells_of_every_scalar_type_end_up_the_same_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    template = named_cells_template(
        tmp_path / "template.xlsx",
        cells={
            "fund_name": ("A1", "placeholder"),
            "fund_nav": ("A2", 0),
            "report_date": ("A3", None),
        },
    )
    values = {"fund_name": "NAE", "fund_nav": 12.5, "report_date": datetime(2025, 1, 2, 9, 30)}
    parity.assert_same(*parity.fill_both(template, values=values))


def test_a_table_on_a_sheet_that_is_not_active_can_still_grow(
    tmp_path: Path, parity: Parity
) -> None:
    """Resizing must not depend on which sheet Excel happens to show."""
    template = table_template(
        tmp_path / "template.xlsx", other_sheets=("Report",), active_sheet="Report"
    )
    rows = [{"fund": f"F{i}", "nav": i} for i in range(6)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))


def test_a_table_name_in_the_wrong_case_is_not_found_by_either_engine(
    tmp_path: Path,
) -> None:
    """Excel's ListObjects() ignores case; openpyxl does not, and both must agree."""
    template = table_template(tmp_path / "template.xlsx")

    for engine in ("openpyxl", "win32com"):
        with pytest.raises(KeyError):
            fill(
                template,
                tmp_path / engine / "filled.xlsx",
                engine=engine,
                tables={"sales": [{"fund": "A", "nav": 1}]},
            )


def test_an_empty_table_with_only_a_blank_row_grows_the_same_way(
    tmp_path: Path, parity: Parity
) -> None:
    """A cleared template row is not content; the table may still grow over it."""
    template = table_template(tmp_path / "template.xlsx", body=((None, None),))
    rows = [{"fund": f"F{i}", "nav": i} for i in range(4)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))


def test_a_header_with_a_line_break_matches_the_data_column_of_the_same_name(
    tmp_path: Path, parity: Parity
) -> None:
    """The file format escapes the line break; COM returns it as it is shown."""
    template = table_template(
        tmp_path / "template.xlsx",
        header=("fund", "NAV_x000a_SEK"),
        header_cells=("fund", "NAV\nSEK"),
    )
    rows = [{"fund": f"F{i}", "NAV\nSEK": i} for i in range(3)]
    parity.assert_same(
        *parity.fill_both(
            template,
            tables={"Sales": rows},
            config=FillXLConfig(column_validation_strategy="strict"),
        )
    )


def test_rich_text_above_a_growing_table_is_preserved_by_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    """A partly bold title is not template furniture either engine may flatten."""
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont

    template = table_template(tmp_path / "template.xlsx", top_row=3)
    workbook = openpyxl.load_workbook(template, rich_text=True)
    workbook["Data"]["A1"] = CellRichText([TextBlock(InlineFont(b=True), "Bold"), " plain"])
    workbook.save(template)
    workbook.close()

    rows = [{"fund": f"F{i}", "nav": i} for i in range(5)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))


def test_a_sorted_table_keeps_its_sort_state_inside_the_body_in_both_engines(
    tmp_path: Path, parity: Parity
) -> None:
    """A sortState outside the table is one of the causes of a repair prompt."""
    from openpyxl import load_workbook
    from openpyxl.worksheet.filters import SortCondition, SortState

    body = tuple((f"F{i}", i) for i in range(3))
    template = table_template(tmp_path / "template.xlsx", body=body)
    workbook = load_workbook(template)
    table = workbook["Data"].tables["Sales"]
    table.sortState = SortState(
        ref="A2:B4", sortCondition=[SortCondition(ref="A2:A4", sortBy="value")]
    )
    workbook.save(template)
    workbook.close()

    rows = [{"fund": f"F{i}", "nav": i} for i in range(7)]
    parity.assert_same(*parity.fill_both(template, tables={"Sales": rows}))
