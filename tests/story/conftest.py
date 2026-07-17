"""Shared fixtures for fillxl story tests.

Creates minimal Excel workbooks programmatically so story tests are
self-contained — no pre-built .xlsx files needed.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest
from openpyxl.worksheet.table import Table, TableColumn  # ty:ignore[unresolved-import]

if TYPE_CHECKING:
    from pathlib import Path


@pytest.fixture
def output_path(tmp_path: Path) -> Path:
    """Return a temporary output path for each test."""
    return tmp_path / "output.xlsx"


# ── Template builders ────────────────────────────────────────────────────────


def _make_named_cell_template(path: Path) -> None:
    """Single named cell 'my_cell' at A1 with value 'hello'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "hello"

    from openpyxl.workbook.defined_name import DefinedName

    wb.defined_names.add(DefinedName("my_cell", attr_text="Sheet1!$A$1"))
    wb.save(path)
    wb.close()


def _make_simple_table_template(path: Path) -> None:
    """Create a table 'SimpleTable' with columns Name + Value, 1 data row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "placeholder"
    ws["B2"] = 0
    table = Table(
        displayName="SimpleTable",
        ref="A1:B2",
        tableColumns=[
            TableColumn(id=1, name="Name"),
            TableColumn(id=2, name="Value"),
        ],
    )
    ws.add_table(table)
    wb.save(path)
    wb.close()


def _make_multi_table_template(path: Path) -> None:
    """Two tables: TableOne + TableTwo on the same sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    # TableOne at A1:B2
    ws["A1"] = "Item"
    ws["B1"] = "Qty"
    ws["A2"] = "x"
    ws["B2"] = 1
    t1 = Table(
        displayName="TableOne",
        ref="A1:B2",
        tableColumns=[TableColumn(id=1, name="Item"), TableColumn(id=2, name="Qty")],
    )
    ws.add_table(t1)

    # TableTwo at D1:E2
    ws["D1"] = "Product"
    ws["E1"] = "Price"
    ws["D2"] = "y"
    ws["E2"] = 99
    t2 = Table(
        displayName="TableTwo",
        ref="D1:E2",
        tableColumns=[TableColumn(id=1, name="Product"), TableColumn(id=2, name="Price")],
    )
    ws.add_table(t2)

    wb.save(path)
    wb.close()


def _make_hidden_sheets_template(path: Path) -> None:
    """Create a workbook with _Config (to hide) and Visible (to keep)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws["A1"] = "Hello"
    wb.create_sheet("_Config").cell(row=1, column=1, value="secret")
    wb.save(path)
    wb.close()


@pytest.fixture(scope="session")
def named_cell_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Workbook with a single named cell 'my_cell'."""
    path = tmp_path_factory.mktemp("story") / "named_cell.xlsx"
    _make_named_cell_template(path)
    return path


@pytest.fixture(scope="session")
def simple_table_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Workbook with table 'SimpleTable' (Name, Value)."""
    path = tmp_path_factory.mktemp("story") / "simple_table.xlsx"
    _make_simple_table_template(path)
    return path


@pytest.fixture(scope="session")
def multi_table_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Workbook with two tables: TableOne + TableTwo."""
    path = tmp_path_factory.mktemp("story") / "multi_table.xlsx"
    _make_multi_table_template(path)
    return path


@pytest.fixture(scope="session")
def hidden_sheets_template(tmp_path_factory: pytest.TempPathFactory) -> Path:
    """Workbook with _Config (to hide) + Visible sheets."""
    path = tmp_path_factory.mktemp("story") / "hidden_sheets.xlsx"
    _make_hidden_sheets_template(path)
    return path
