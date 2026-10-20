"""Chapter 3 — Filling Excel Tables.
==================================

An **Excel table** (Insert → Table, or ``Ctrl+T``) is a named range of
cells with headers and structured references.  Tables are the heart of
fillxl: you pass tabular data from Python and fillxl writes it into the
matching Excel table, automatically **resizing** the table to fit.

When your data has more rows than the template table, the table
**grows** downward.  When it has fewer, the table **shrinks**.  Existing
styles, formatting, calculated-column formulas, and conditional
formatting are all preserved and extended to new rows.

Data can come from any of five formats — ``list[dict]``,
``dict[str, list]``, pandas DataFrame, polars DataFrame, or duckdb
relation (covered in Chapter 4).  This chapter focuses on the core
table-filling workflow, edge cases, and style preservation.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Any, cast

import openpyxl
import pytest

from fillxl import FillXLConfig, fill

if TYPE_CHECKING:
    from pathlib import Path

# ── Helpers ──────────────────────────────────────────────────────────────────


def _make_table_template(
    path: Path,
    columns: list[str],
    rows: int = 1,
    table_name: str = "MyTable",
) -> None:
    """Create a workbook with a single table."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    from openpyxl.worksheet.table import Table, TableColumn

    for ci, col_name in enumerate(columns, 1):
        ws.cell(row=1, column=ci, value=col_name)
    for ri in range(rows):
        for ci in range(len(columns)):
            ws.cell(row=ri + 2, column=ci + 1, value=f"r{ri}c{ci}")

    last_col_letter = openpyxl.utils.get_column_letter(len(columns))
    ref = f"A1:{last_col_letter}{rows + 1}"
    table = Table(
        displayName=table_name,
        ref=ref,
        tableColumns=[TableColumn(id=i + 1, name=c) for i, c in enumerate(columns)],
    )
    ws.add_table(table)
    wb.save(path)
    wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 3.1 Basic table filling
# ═══════════════════════════════════════════════════════════════════════════════


class TestFillingTables:
    """Basic table-filling: writing data, growing, shrinking, multiple tables."""

    def test_data_is_written_into_the_table_cells(
        self,
        output_path: Path,
        simple_table_template: Path,
    ) -> None:
        """Values from your data appear in the correct cells."""
        data = [{"Name": "Alice", "Value": 30}, {"Name": "Bob", "Value": 25}]

        fill(
            template=simple_table_template,
            output=output_path,
            tables={"SimpleTable": data},
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Data"]
            assert ws["A1"].value == "Name"  # header unchanged
            assert ws["A2"].value == "Alice"
            assert ws["B2"].value == 30
            assert ws["A3"].value == "Bob"
            assert ws["B3"].value == 25
        finally:
            wb.close()

    def test_the_table_grows_when_data_has_more_rows_than_the_template(
        self,
        output_path: Path,
    ) -> None:
        """If you supply 10 rows but the template table has only 1, the table expands."""
        _make_table_template(
            output_path.with_name("grow_tmpl.xlsx"),
            columns=["X", "Y"],
            rows=1,
            table_name="GrowTable",
        )

        data = [{"X": f"item{i}", "Y": i} for i in range(1, 11)]

        fill(
            template=output_path.with_name("grow_tmpl.xlsx"),
            output=output_path,
            tables={"GrowTable": data},
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Data"]
            assert ws["A2"].value == "item1"
            assert ws["A11"].value == "item10"
            # Row 12 should be empty (table = header + 10 data rows = 11 rows)
            assert ws["A12"].value is None
        finally:
            wb.close()

    def test_the_table_shrinks_when_data_has_fewer_rows_than_the_template(
        self,
        output_path: Path,
    ) -> None:
        """If you supply 2 rows but the template has 5, the table contracts."""
        _make_table_template(
            output_path.with_name("shrink_tmpl.xlsx"),
            columns=["Col"],
            rows=5,
            table_name="ShrinkTable",
        )

        data = [{"Col": "a"}, {"Col": "b"}]

        fill(
            template=output_path.with_name("shrink_tmpl.xlsx"),
            output=output_path,
            tables={"ShrinkTable": data},
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Data"]
            assert ws["A2"].value == "a"
            assert ws["A3"].value == "b"
            # Old row 4 should be cleared
            assert ws["A4"].value is None
            assert ws["A6"].value is None
        finally:
            wb.close()

    def test_multiple_tables_can_be_filled_in_one_call(
        self,
        output_path: Path,
        multi_table_template: Path,
    ) -> None:
        """Filling two tables at once updates both independently."""
        fill(
            template=multi_table_template,
            output=output_path,
            tables={
                "TableOne": [{"Item": "Widget", "Qty": 10}],
                "TableTwo": [{"Product": "Gadget", "Price": 199}],
            },
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Data"]
            assert ws["A2"].value == "Widget"
            assert ws["B2"].value == 10
            assert ws["D2"].value == "Gadget"
            assert ws["E2"].value == 199
        finally:
            wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 3.2 Edge cases
# ═══════════════════════════════════════════════════════════════════════════════


class TestTableEdgeCases:
    """Zero-row data, missing table names, non-tabular data, overwrite protection."""

    def test_filling_with_zero_data_rows_leaves_only_the_header(
        self,
        output_path: Path,
        simple_table_template: Path,
    ) -> None:
        """An empty dataset clears all data rows but keeps the header."""
        fill(
            template=simple_table_template,
            output=output_path,
            tables={"SimpleTable": []},
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Data"]
            assert ws["A1"].value == "Name"  # header intact
            assert ws["A2"].value is None  # data cleared
        finally:
            wb.close()

    def test_a_keyerror_is_raised_when_the_table_name_is_not_in_the_workbook(
        self,
        output_path: Path,
        simple_table_template: Path,
    ) -> None:
        """Referencing a table that doesn't exist raises KeyError."""
        with pytest.raises(KeyError):
            fill(
                template=simple_table_template,
                output=output_path,
                tables={"NonExistentTable": [{"X": 1}]},
            )

    def test_a_typeerror_is_raised_for_data_that_is_not_tabular(
        self,
        output_path: Path,
        simple_table_template: Path,
    ) -> None:
        """Passing a plain string or number where tabular data is expected fails."""
        with pytest.raises(TypeError, match="Unsupported data type"):
            fill(
                template=simple_table_template,
                output=output_path,
                tables={"SimpleTable": cast(Any, "not tabular data")},
            )

    def test_table_growth_is_blocked_when_it_would_overwrite_existing_content(
        self,
        output_path: Path,
    ) -> None:
        """If expanding the table would overwrite non-empty cells, an error is raised."""
        _make_table_template(
            output_path.with_name("blocked_tmpl.xlsx"),
            columns=["Col"],
            rows=1,
            table_name="BlockedTable",
        )

        # Place content directly below the table
        wb = openpyxl.load_workbook(output_path.with_name("blocked_tmpl.xlsx"))
        ws = wb["Data"]
        ws["A5"] = "DO NOT OVERWRITE"
        wb.save(output_path.with_name("blocked_tmpl.xlsx"))
        wb.close()

        data = [{"Col": f"row{i}"} for i in range(1, 10)]

        with pytest.raises(ValueError, match="cannot expand"):
            fill(
                template=output_path.with_name("blocked_tmpl.xlsx"),
                output=output_path,
                tables={"BlockedTable": data},
            )


# ═══════════════════════════════════════════════════════════════════════════════
# 3.3 Style and formatting preservation
# ═══════════════════════════════════════════════════════════════════════════════


class TestTableStylePreservation:
    """New rows inherit styling; style copying can be disabled."""

    def test_new_rows_inherit_the_style_of_the_first_data_row(
        self,
        output_path: Path,
    ) -> None:
        """When copy_table_styles is True (default), new rows copy row-1 styling."""
        _make_table_template(
            output_path.with_name("styled_tmpl.xlsx"),
            columns=["Name", "Score"],
            rows=1,
            table_name="StyledTable",
        )

        # Apply a custom style to the first data row
        wb = openpyxl.load_workbook(output_path.with_name("styled_tmpl.xlsx"))
        ws = wb["Data"]
        from openpyxl.styles import Font, PatternFill

        ws["A2"].font = Font(name="Calibri", size=12, bold=True)
        ws["A2"].fill = PatternFill(start_color="FFAAAA", end_color="FFAAAA", fill_type="solid")
        wb.save(output_path.with_name("styled_tmpl.xlsx"))
        wb.close()

        data = [{"Name": "Alice", "Score": 95}, {"Name": "Bob", "Score": 87}]

        fill(
            template=output_path.with_name("styled_tmpl.xlsx"),
            output=output_path,
            tables={"StyledTable": data},
        )

        wb2 = openpyxl.load_workbook(output_path)
        try:
            ws2 = wb2["Data"]
            # Row 3 should have same font as row 2
            assert ws2["A3"].font.bold is True
            assert ws2["A3"].font.name == "Calibri"
        finally:
            wb2.close()

    def test_style_copying_can_be_disabled_with_copy_table_styles_false(
        self,
        output_path: Path,
    ) -> None:
        """Set copy_table_styles=False to skip style inheritance."""
        _make_table_template(
            output_path.with_name("nostyle_tmpl.xlsx"),
            columns=["X"],
            rows=1,
            table_name="NoStyleTable",
        )

        data = [{"X": "a"}, {"X": "b"}, {"X": "c"}]

        fill(
            template=output_path.with_name("nostyle_tmpl.xlsx"),
            output=output_path,
            tables={"NoStyleTable": data},
            config=FillXLConfig(copy_table_styles=False),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Data"]
            # Data should be there
            assert ws["A2"].value == "a"
            assert ws["A4"].value == "c"
        finally:
            wb.close()


# ══════════════════════════════════════════════════════════════════════════════
# 3.5 Large data
# ══════════════════════════════════════════════════════════════════════════════


class TestLargeData:
    """Filling with large amounts of data does not crash."""

    @pytest.mark.slow
    def test_large_data_does_not_crash(
        self,
        output_path: Path,
    ) -> None:
        """Filling a table with 1000 rows completes without error."""
        from openpyxl.worksheet.table import Table, TableColumn

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "ID"
        ws["B1"] = "Name"
        ws["C1"] = "Value"
        ws["A2"] = 0
        ws["B2"] = "placeholder"
        ws["C2"] = 0.0

        table = Table(
            displayName="BigTable",
            ref="A1:C2",
            tableColumns=[
                TableColumn(id=1, name="ID"),
                TableColumn(id=2, name="Name"),
                TableColumn(id=3, name="Value"),
            ],
        )
        ws.add_table(table)

        tmpl = output_path.with_name("big_tmpl.xlsx")
        wb.save(tmpl)
        wb.close()

        # Generate 1000 rows
        data = [{"ID": i, "Name": f"Item_{i}", "Value": i * 1.5} for i in range(1000)]

        # Should complete without error
        fill(template=tmpl, output=output_path, tables={"BigTable": data})

        wb2 = openpyxl.load_workbook(output_path)
        try:
            ws2 = wb2["Data"]
            assert ws2["A2"].value == 0
            assert ws2["A1001"].value == 999
            assert ws2["C1001"].value == 999 * 1.5
        finally:
            wb2.close()
