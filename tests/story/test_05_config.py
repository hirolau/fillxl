"""Chapter 5 — Configuring fillxl.
================================

:class:`~fillxl.FillXLConfig` gives you fine-grained control over how
fillxl behaves.  Every option has a sensible default, so you only need
to configure what you want to change.

Configuration is passed as the ``config`` keyword argument to :func:`fillxl.fill`.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest

from fillxl import FillXLConfig, fill

if TYPE_CHECKING:
    from pathlib import Path

    from fillxl._validation import ColumnValidationStrategy

# ── Helpers ──────────────────────────────────────────────────────────────────


def _make_simple_template(path: Path) -> None:
    """Create a minimal workbook with a table + named cell."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"

    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table, TableColumn  # ty:ignore[unresolved-import]

    ws["A1"] = "Item"
    ws["B1"] = "Qty"
    ws["A2"] = "old"
    ws["B2"] = 0
    ws["C1"] = "Old Title"
    wb.defined_names.add(DefinedName("title", attr_text="Data!$C$1"))

    table = Table(
        displayName="Items",
        ref="A1:B2",
        tableColumns=[
            TableColumn(id=1, name="Item"),
            TableColumn(id=2, name="Qty"),
        ],
    )
    ws.add_table(table)
    wb.save(path)
    wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 5.1 Engine selection
# ═══════════════════════════════════════════════════════════════════════════════


class TestEngineSelection:
    """Engine defaults to 'auto' (openpyxl); 'win32com' accepted on Windows."""

    def test_the_default_engine_is_auto_which_uses_openpyxl(
        self,
        output_path: Path,
    ) -> None:
        """With engine='auto' (the default), openpyxl is used."""
        _make_simple_template(output_path.with_name("tmpl.xlsx"))

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            values={"title": "Default Engine"},
            engine="auto",
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["Data"]["C1"].value == "Default Engine"
        finally:
            wb.close()

    def test_engine_can_be_explicitly_set_to_openpyxl(
        self,
        output_path: Path,
    ) -> None:
        """engine='openpyxl' explicitly selects the openpyxl backend."""
        _make_simple_template(output_path.with_name("tmpl.xlsx"))

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            values={"title": "OpenPyXL"},
            engine="openpyxl",
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["Data"]["C1"].value == "OpenPyXL"
        finally:
            wb.close()

    def test_engine_win32com_is_accepted_on_windows(self) -> None:
        """The 'win32com' engine name is accepted by get_backend()."""
        from fillxl._backend import get_backend

        backend = get_backend("win32com")
        assert backend is not None

    def test_an_invalid_engine_name_raises_valueerror(self) -> None:
        """'xlwings' or any unknown engine name is rejected by get_backend()."""
        from fillxl._backend import get_backend

        with pytest.raises(ValueError, match="Unknown engine"):
            get_backend("xlwings")  # type: ignore[arg-type]  # ty:ignore[invalid-argument-type]


# ═══════════════════════════════════════════════════════════════════════════════
# 5.2 Table options
# ═══════════════════════════════════════════════════════════════════════════════


class TestTableOptions:
    """extend_tables, copy_table_styles, extend_conditional_formatting defaults."""

    def test_extend_tables_is_enabled_by_default(self) -> None:
        """By default tables resize to fit your data."""
        cfg = FillXLConfig()
        assert cfg.extend_tables is True

    def test_copy_table_styles_is_enabled_by_default(self) -> None:
        """By default new rows inherit the first row's formatting."""
        cfg = FillXLConfig()
        assert cfg.copy_table_styles is True

    def test_extend_conditional_formatting_is_enabled_by_default(self) -> None:
        """By default CF rules stretch to cover new rows."""
        cfg = FillXLConfig()
        assert cfg.extend_conditional_formatting is True


# ═══════════════════════════════════════════════════════════════════════════════
# 5.3 Sheet management options
# ═══════════════════════════════════════════════════════════════════════════════


class TestSheetManagementOptions:
    """hide_underscore_sheets, refresh_pivot_tables, delete_sheets, active_sheet."""

    def test_hide_underscore_sheets_is_enabled_by_default(self) -> None:
        """Sheets starting with '_' are hidden by default."""
        cfg = FillXLConfig()
        assert cfg.hide_underscore_sheets is True

    def test_refresh_pivot_tables_is_enabled_by_default(self) -> None:
        """Pivot caches are marked for refresh by default."""
        cfg = FillXLConfig()
        assert cfg.refresh_pivot_tables is True

    def test_delete_sheets_is_none_by_default(self) -> None:
        """No sheets are deleted unless you opt in."""
        cfg = FillXLConfig()
        assert cfg.delete_sheets is None

    def test_active_sheet_is_none_by_default(self) -> None:
        """The active sheet is not changed unless you specify one."""
        cfg = FillXLConfig()
        assert cfg.active_sheet is None


# ═══════════════════════════════════════════════════════════════════════════════
# 5.4 Column validation strategy
# ═══════════════════════════════════════════════════════════════════════════════


class TestColumnValidationStrategyOption:
    """Default is 'all_but_formulas'; six valid names; invalid names rejected."""

    def test_the_default_strategy_is_all_but_formulas(self) -> None:
        """By default, all non-formula table columns must be present in data."""
        cfg = FillXLConfig()
        assert cfg.column_validation_strategy == "all_but_formulas"

    @pytest.mark.parametrize(
        "strategy",
        [
            "strict",
            "all_but_formulas",
            "all_but_formulas_strict",
            "none",
            "drop",
            "warn",
        ],
        ids=["strict", "all_but_formulas", "all_but_formulas_strict", "none", "drop", "warn"],
    )
    def test_strategy_name_is_accepted(self, strategy: ColumnValidationStrategy) -> None:
        """Each of the six strategy names passes config validation."""
        cfg = FillXLConfig(column_validation_strategy=strategy)
        assert cfg.column_validation_strategy == strategy

    def test_an_invalid_strategy_name_raises_valueerror(self) -> None:
        """'bogus' or any unknown strategy is rejected."""
        with pytest.raises(ValueError, match="column_validation_strategy"):
            FillXLConfig(column_validation_strategy="bogus")  # type: ignore[arg-type]  # ty:ignore[invalid-argument-type]


# ═══════════════════════════════════════════════════════════════════════════════
# 5.5 Config is frozen (immutable)
# ═══════════════════════════════════════════════════════════════════════════════


class TestConfigIsImmutable:
    """FillXLConfig is frozen — attribute assignment raises an exception."""

    def test_config_instances_are_frozen_and_cannot_be_changed(self) -> None:
        """Attempting to set an attribute on a config raises an exception."""
        import attrs

        cfg = FillXLConfig()
        with pytest.raises(attrs.exceptions.FrozenInstanceError):
            cfg.refresh_pivot_tables = False  # type: ignore[misc]  # ty:ignore[invalid-assignment]
