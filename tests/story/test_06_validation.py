"""Chapter 6 — Column Validation.
===============================

Sometimes the columns in your data don't exactly match the columns in the
Excel table.  fillxl gives you **six strategies** for handling this:

=============  ===========================================================
Strategy        Behaviour
=============  ===========================================================
``strict``     All columns must match 1-to-1.  Extra or missing → error.
``all_but_formulas`` (default)  All non-formula table columns required.
                Extra data columns allowed.  Formula columns ignored.
``all_but_formulas_strict``  Like above, but extra data columns → error.
``none``       No validation at all — just write what you have.
``drop``       Silently drop extra data columns, skip missing ones.
``warn``       Log warnings but proceed with all data columns.
=============  ===========================================================

Set the strategy via ``FillXLConfig(column_validation_strategy=...)``.
"""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

import openpyxl
import pytest

from fillxl import ColumnValidationStrategy, FillXLConfig, fill
from fillxl._validation import validate_columns

if TYPE_CHECKING:
    from pathlib import Path

# ── Helpers ──────────────────────────────────────────────────────────────────


def _make_table_template(path: Path, columns: list[str], table_name: str = "Data") -> None:
    """Create a workbook with a table having the given columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    from openpyxl.worksheet.table import Table, TableColumn

    for ci, col_name in enumerate(columns, 1):
        ws.cell(row=1, column=ci, value=col_name)
        ws.cell(row=2, column=ci, value=f"dummy_{col_name}")

    last_col = openpyxl.utils.get_column_letter(len(columns))
    table = Table(
        displayName=table_name,
        ref=f"A1:{last_col}2",
        tableColumns=[TableColumn(id=i + 1, name=c) for i, c in enumerate(columns)],
    )
    ws.add_table(table)
    wb.save(path)
    wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 6.1 The default strategy: all_but_formulas
# ═══════════════════════════════════════════════════════════════════════════════


class TestAllButFormulasStrategy:
    """Default strategy: non-formula columns required, extra columns allowed."""

    def test_missing_a_non_formula_column_raises_valueerror(
        self,
        output_path: Path,
    ) -> None:
        """If a non-formula table column is missing from the data, fill() raises."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["Name", "Age", "City"],
        )

        # Missing 'City'
        data = [{"Name": "Alice", "Age": 30}]

        with pytest.raises(ValueError, match="no data column"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
                config=FillXLConfig(column_validation_strategy="all_but_formulas"),
            )

    def test_extra_columns_in_the_data_are_allowed(
        self,
        output_path: Path,
    ) -> None:
        """Extra data columns beyond what the table expects are silently accepted."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["Name", "Age"],
        )

        data = [{"Name": "Alice", "Age": 30, "ExtraField": "ignored"}]

        # Should not raise
        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            tables={"Data": data},
            config=FillXLConfig(column_validation_strategy="all_but_formulas"),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Sheet1"]
            assert ws["A2"].value == "Alice"
        finally:
            wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 6.2 Strict strategy
# ═══════════════════════════════════════════════════════════════════════════════


class TestStrictStrategy:
    """Exact 1-to-1 column matching; extra or missing columns raise."""

    def test_exact_column_match_passes_silently(
        self,
        output_path: Path,
    ) -> None:
        """When columns match exactly, strict mode is happy."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["A", "B"],
        )

        data = [{"A": 1, "B": 2}]

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            tables={"Data": data},
            config=FillXLConfig(column_validation_strategy="strict"),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["Sheet1"]["A2"].value == 1
        finally:
            wb.close()

    def test_extra_data_column_raises_valueerror(
        self,
        output_path: Path,
    ) -> None:
        """An extra column in data that's not in the table → error."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["A"],
        )

        data = [{"A": 1, "B": 2}]

        with pytest.raises(ValueError, match="Column mismatch"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
                config=FillXLConfig(column_validation_strategy="strict"),
            )

    def test_missing_table_column_raises_valueerror(
        self,
        output_path: Path,
    ) -> None:
        """A table column missing from the data → error."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["A", "B"],
        )

        data = [{"A": 1}]

        with pytest.raises(ValueError, match="Column mismatch"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
                config=FillXLConfig(column_validation_strategy="strict"),
            )


# ═══════════════════════════════════════════════════════════════════════════════
# 6.3 Drop and warn strategies
# ═══════════════════════════════════════════════════════════════════════════════


class TestDropAndWarnStrategies:
    """Permissive strategies: drop removes extras, warn logs, none skips all."""

    def test_drop_strategy_silently_removes_extra_columns(
        self,
        output_path: Path,
    ) -> None:
        """Extra data columns are dropped; missing table columns are skipped."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["Name"],
        )

        data = [{"Name": "Alice", "Extra": "gone"}]

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            tables={"Data": data},
            config=FillXLConfig(column_validation_strategy="drop"),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["Sheet1"]["A2"].value == "Alice"
        finally:
            wb.close()

    def test_warn_strategy_logs_warnings_and_proceeds(
        self,
        output_path: Path,
    ) -> None:
        """Mismatches are logged at WARNING level but fill continues."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["Name"],
        )

        data = [{"Name": "Alice", "Extra": "warned"}]

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            tables={"Data": data},
            config=FillXLConfig(column_validation_strategy="warn"),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["Sheet1"]["A2"].value == "Alice"
        finally:
            wb.close()

    def test_none_strategy_skips_validation_entirely(
        self,
        output_path: Path,
    ) -> None:
        """No checks at all — just write whatever columns match."""
        _make_table_template(
            output_path.with_name("tmpl.xlsx"),
            columns=["Name"],
        )

        data = [{"Name": "Alice", "Extra": "ignored", "More": "also"}]

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            tables={"Data": data},
            config=FillXLConfig(column_validation_strategy="none"),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["Sheet1"]["A2"].value == "Alice"
        finally:
            wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 6.4 Duplicate column names
# ═══════════════════════════════════════════════════════════════════════════════


class TestDuplicateColumnNames:
    """Ambiguous column names are rejected before anything is written."""

    def test_two_columns_with_the_same_label_are_rejected(
        self,
        output_path: Path,
    ) -> None:
        """A DataFrame with two "amount" columns cannot be filled unambiguously."""
        pd = pytest.importorskip("pandas")
        _make_table_template(output_path.with_name("tmpl.xlsx"), columns=["Region", "Amount"])

        data = pd.DataFrame([["North", 10, 20]], columns=["Region", "amount", "amount"])

        with pytest.raises(ValueError, match=r"Duplicate column names.*'amount'"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
            )

    def test_labels_that_normalize_to_the_same_name_are_rejected(
        self,
        output_path: Path,
    ) -> None:
        """A numeric header ``1`` and a text header ``"1"`` both become column "1"."""
        _make_table_template(output_path.with_name("tmpl.xlsx"), columns=["Region", "1"])

        # Non-string mapping keys are supported, so this reaches the adapter.
        data: Any = {1: ["North"], "1": [10]}

        with pytest.raises(ValueError, match=r"Duplicate column names.*'1'"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
            )

    def test_no_output_is_written_when_columns_are_ambiguous(
        self,
        output_path: Path,
    ) -> None:
        """The fill stops before writing, so no column is silently dropped."""
        _make_table_template(output_path.with_name("tmpl.xlsx"), columns=["Region", "1"])
        output_path.write_bytes(b"previous")

        data: Any = {1: ["North"], "1": [10]}

        with pytest.raises(ValueError, match="Duplicate column names"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
            )

        assert output_path.read_bytes() == b"previous"

    @pytest.mark.parametrize("strategy", ["strict", "all_but_formulas", "warn", "drop", "none"])
    def test_the_check_does_not_depend_on_the_validation_strategy(
        self,
        output_path: Path,
        strategy: ColumnValidationStrategy,
    ) -> None:
        """Duplicate names are ambiguous even where validation is switched off."""
        _make_table_template(output_path.with_name("tmpl.xlsx"), columns=["Region", "1"])

        data: Any = {1: ["North"], "1": [10]}

        with pytest.raises(ValueError, match="Duplicate column names"):
            fill(
                template=output_path.with_name("tmpl.xlsx"),
                output=output_path,
                tables={"Data": data},
                config=FillXLConfig(column_validation_strategy=strategy),
            )

    def test_unique_column_names_still_fill(
        self,
        output_path: Path,
    ) -> None:
        """Distinct labels keep the existing behaviour."""
        _make_table_template(output_path.with_name("tmpl.xlsx"), columns=["Region", "Amount"])

        fill(
            template=output_path.with_name("tmpl.xlsx"),
            output=output_path,
            tables={"Data": [{"Region": "North", "Amount": 10}]},
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            ws = wb["Sheet1"]
            assert ws["A2"].value == "North"
            assert ws["B2"].value == 10
        finally:
            wb.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 6.5 The validate_columns function (unit level)
# ═══════════════════════════════════════════════════════════════════════════════


class TestValidateColumnsFunction:
    """Unit tests for the validate_columns internal function.

    The function validates and reports; it returns ``None`` because the
    backends match columns by name and never need a filtered list.
    """

    def test_none_strategy_accepts_any_column_combination(self) -> None:
        """'none' accepts whatever you pass in."""
        validate_columns(
            table_columns=["A", "B"],
            data_columns=["X", "Y", "Z"],
            strategy="none",
        )

    def test_warn_strategy_accepts_mismatched_columns(self) -> None:
        """'warn' accepts extra data columns."""
        validate_columns(
            table_columns=["A"],
            data_columns=["A", "B", "C"],
            strategy="warn",
        )

    def test_drop_strategy_accepts_extra_and_missing_columns(self) -> None:
        """'drop' accepts both extra data columns and missing table columns."""
        validate_columns(
            table_columns=["A", "B"],
            data_columns=["B", "C"],
            strategy="drop",
        )

    def test_strict_raises_when_columns_differ(self) -> None:
        """'strict' raises ValueError on any mismatch."""
        with pytest.raises(ValueError, match="Column mismatch"):
            validate_columns(
                table_columns=["A", "B"],
                data_columns=["A", "C"],
                strategy="strict",
            )

    def test_all_but_formulas_ignores_formula_columns(self) -> None:
        """Formula columns are excluded from the comparison."""
        # 'B' is a formula column — it's OK if missing from data
        # No error expected since all non-formula cols match
        validate_columns(
            table_columns=["A", "B", "C"],
            data_columns=["A", "C"],
            strategy="all_but_formulas",
            formula_columns=["B"],
        )

    def test_all_but_formulas_strict_rejects_extra_data_columns(self) -> None:
        """Extra data columns not in the table → error in strict mode."""
        with pytest.raises(ValueError, match="extra columns"):
            validate_columns(
                table_columns=["A", "B"],
                data_columns=["A", "B", "C"],
                strategy="all_but_formulas_strict",
                formula_columns=[],
            )
