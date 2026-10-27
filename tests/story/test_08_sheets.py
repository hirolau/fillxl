"""Chapter 8 — Sheet Management.
==============================

Beyond filling cells and tables, fillxl can also manage the **sheets**
in your workbook:

* **Hide** sheets whose names start with ``_`` (like ``_Config``).
* **Delete** specific sheets entirely.
* **Set the active sheet** — the one Excel displays on open.

These are all controlled via :class:`~fillxl.FillXLConfig`.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import openpyxl
import pytest

from fillxl import FillXLConfig, IgnoredDataWarning, fill

if TYPE_CHECKING:
    from pathlib import Path

    import pytest

# ═══════════════════════════════════════════════════════════════════════════════
# 8.1 Hiding underscore-prefixed sheets
# ═══════════════════════════════════════════════════════════════════════════════


class TestHidingUnderscoreSheets:
    """Sheets starting with '_' are hidden by default; can be disabled."""

    def test_a_sheet_starting_with_underscore_is_hidden_after_fill(
        self,
        output_path: Path,
        hidden_sheets_template: Path,
    ) -> None:
        """_Config becomes hidden; Visible stays visible."""
        fill(template=hidden_sheets_template, output=output_path)

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["_Config"].sheet_state == "hidden"
            assert wb["Visible"].sheet_state == "visible"
        finally:
            wb.close()

    def test_underscore_sheets_remain_visible_when_hide_flag_is_false(
        self,
        output_path: Path,
        hidden_sheets_template: Path,
    ) -> None:
        """Set hide_underscore_sheets=False to keep them visible."""
        fill(
            template=hidden_sheets_template,
            output=output_path,
            config=FillXLConfig(hide_underscore_sheets=False),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb["_Config"].sheet_state == "visible"
        finally:
            wb.close()

    def test_hide_underscore_sheets_is_enabled_by_default(self) -> None:
        """The default config hides underscore sheets."""
        cfg = FillXLConfig()
        assert cfg.hide_underscore_sheets is True


# ═══════════════════════════════════════════════════════════════════════════════
# 8.2 Deleting sheets
# ═══════════════════════════════════════════════════════════════════════════════


class TestDeletingSheets:
    """Sheets listed in delete_sheets are removed; missing names warn."""

    def test_a_sheet_listed_in_delete_sheets_is_removed(
        self,
        output_path: Path,
        hidden_sheets_template: Path,
    ) -> None:
        """Specify delete_sheets=['_Config'] and it's gone from the output."""
        fill(
            template=hidden_sheets_template,
            output=output_path,
            config=FillXLConfig(delete_sheets=["_Config"]),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert "_Config" not in wb.sheetnames
            assert "Visible" in wb.sheetnames
        finally:
            wb.close()

    def test_deleting_a_nonexistent_sheet_warns_not_errors(
        self,
        output_path: Path,
        hidden_sheets_template: Path,
    ) -> None:
        """If you ask to delete a sheet that doesn't exist, fill() warns and continues."""
        with pytest.warns(IgnoredDataWarning, match="not found"):
            fill(
                template=hidden_sheets_template,
                output=output_path,
                config=FillXLConfig(delete_sheets=["GhostSheet"]),
            )


# ═══════════════════════════════════════════════════════════════════════════════
# 8.3 Setting the active sheet
# ═══════════════════════════════════════════════════════════════════════════════


class TestActiveSheet:
    """The active sheet can be set; invalid names warn instead of erroring."""

    def test_the_active_sheet_is_set_to_the_specified_name(
        self,
        output_path: Path,
        hidden_sheets_template: Path,
    ) -> None:
        """Set active_sheet='Visible' and that sheet becomes active."""
        fill(
            template=hidden_sheets_template,
            output=output_path,
            config=FillXLConfig(active_sheet="Visible"),
        )

        wb = openpyxl.load_workbook(output_path)
        try:
            assert wb.active.title == "Visible"
        finally:
            wb.close()

    def test_setting_active_sheet_to_a_nonexistent_name_warns(
        self,
        output_path: Path,
        hidden_sheets_template: Path,
    ) -> None:
        """A bad active_sheet name does not crash — just a warning."""
        with pytest.warns(IgnoredDataWarning, match="not found"):
            fill(
                template=hidden_sheets_template,
                output=output_path,
                config=FillXLConfig(active_sheet="NoSuchSheet"),
            )


# ═══════════════════════════════════════════════════════════════════════════════
# 8.4 Sheet protection
# ═══════════════════════════════════════════════════════════════════════════════


class TestProtectedSheets:
    """openpyxl does not enforce sheet protection at the library level, so
    fillxl writes to protected sheets without error.  This test documents
    that behaviour — protection is only enforced when the file is opened
    in Excel.
    """

    def test_protected_sheets_are_written_without_error(
        self,
        output_path: Path,
    ) -> None:
        """Writing to a protected sheet does NOT raise — openpyxl doesn't enforce locks."""
        import openpyxl
        from openpyxl.workbook.defined_name import DefinedName

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Protected"
        ws["A1"] = "secret"
        ws.protection.sheet = True
        ws.protection.password = ""

        wb.defined_names.add(DefinedName("secret_cell", attr_text="Protected!$A$1"))

        tmpl = output_path.with_name("protected.xlsx")
        wb.save(tmpl)
        wb.close()

        # openpyxl does not enforce protection at the library level — this
        # should succeed without error.  Protection is only honoured by Excel.
        fill(template=tmpl, output=output_path, values={"secret_cell": "new_value"})

        wb2 = openpyxl.load_workbook(output_path)
        try:
            assert wb2["Protected"]["A1"].value == "new_value"
            # The sheet should still be marked as protected in the XML
            assert wb2["Protected"].protection.sheet is True
        finally:
            wb2.close()
