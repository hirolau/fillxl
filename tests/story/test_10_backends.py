"""Chapter 10 — Under the Hood: Backends.
=======================================

fillxl uses a **backend** abstraction to talk to Excel files.  There are
two backends:

* **openpyxl** (default, cross-platform) — reads/writes ``.xlsx`` files
  directly using the openpyxl library.  No Excel installation needed.
* **win32com** (Windows only) — uses Microsoft Excel via COM automation.
  Preserves all chart types but requires Excel to be installed.

The :func:`~fillxl._backend.get_backend` factory selects the right backend
based on the ``engine`` config option.  Each backend implements the
:class:`~fillxl._backend.Backend` protocol.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

import openpyxl
import pytest

from fillxl._backend import Backend, get_backend
from fillxl._types import Engine

# ═══════════════════════════════════════════════════════════════════════════════
# 10.1 The Backend factory
# ═══════════════════════════════════════════════════════════════════════════════


class TestBackendFactory:
    """get_backend() maps engine names to backend implementations."""

    def test_get_backend_returns_openpyxl_backend_for_auto(
        self,
    ) -> None:
        """'auto' (the default) maps to openpyxl."""
        from fillxl._backends._openpyxl import OpenpyxlBackend

        backend = get_backend("auto")
        assert isinstance(backend, OpenpyxlBackend)

    def test_get_backend_returns_openpyxl_backend_for_openpyxl(
        self,
    ) -> None:
        """Explicit 'openpyxl' also gives OpenpyxlBackend."""
        from fillxl._backends._openpyxl import OpenpyxlBackend

        backend = get_backend("openpyxl")
        assert isinstance(backend, OpenpyxlBackend)

    def test_get_backend_returns_win32com_backend_for_win32com(
        self,
    ) -> None:
        """'win32com' gives a Win32ComBackend (no Excel launched yet)."""
        from fillxl._backends._win32com import Win32ComBackend

        backend = get_backend("win32com")
        assert isinstance(backend, Win32ComBackend)

    def test_get_backend_raises_valueerror_for_unknown_engine(
        self,
    ) -> None:
        """An unknown engine name raises immediately."""
        with pytest.raises(ValueError, match="Unknown engine"):
            # Intentionally passing invalid engine name to test error handling
            get_backend(cast(Any, "xlwings"))

    @pytest.mark.parametrize("engine", ["auto", "openpyxl", "win32com"])
    def test_known_engine_name_returns_a_valid_backend(self, engine: str) -> None:
        """Every known engine name returns a backend implementing the protocol."""
        backend = get_backend(cast(Engine, engine))
        assert isinstance(backend, Backend)


# ═══════════════════════════════════════════════════════════════════════════════
# 10.2 OpenpyxlBackend guards
# ═══════════════════════════════════════════════════════════════════════════════


class TestOpenpyxlBackendGuards:
    """Operations before load() raise RuntimeError; close() is always safe."""

    def test_operations_raise_runtime_error_if_no_workbook_is_loaded(
        self,
    ) -> None:
        """Calling fill_named_cell or save before load() raises RuntimeError.
        close() is safe to call anytime — it's a no-op when nothing is loaded.
        """
        from fillxl._backends._openpyxl import OpenpyxlBackend
        from fillxl._config import FillXLConfig

        backend = OpenpyxlBackend()

        with pytest.raises(RuntimeError, match="No workbook loaded"):
            backend.fill_named_cell("cell", "value", FillXLConfig())

        with pytest.raises(RuntimeError, match="No workbook loaded"):
            backend.save(Path("out.xlsx"))

        # close() is safe — doesn't raise when workbook is None
        backend.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 10.4 Backend reuse
# ═══════════════════════════════════════════════════════════════════════════════


class TestBackendReuse:
    """A backend instance can be reused across multiple load-fill-save cycles."""

    def test_openpyxl_backend_can_be_reused_for_multiple_cycles(
        self,
        output_path: Path,
    ) -> None:
        """Load → fill → save → close can be repeated with the same instance."""
        from fillxl._backends._openpyxl import OpenpyxlBackend
        from fillxl._tables import populate_table

        backend = OpenpyxlBackend()

        # Cycle 1
        path1 = output_path.with_name("cycle1.xlsx")
        _make_simple_table(path1, "T1")
        backend.load(path1)
        assert backend._workbook is not None
        populate_table(backend._workbook, "T1", [{"Name": "first", "Value": 1}])
        backend.save(path1)
        backend.close()

        # Verify cycle 1
        wb1 = openpyxl.load_workbook(path1)
        try:
            assert wb1["Sheet1"]["A2"].value == "first"
            assert wb1["Sheet1"]["B2"].value == 1
        finally:
            wb1.close()

        # Cycle 2
        path2 = output_path.with_name("cycle2.xlsx")
        _make_simple_table(path2, "T2")
        backend.load(path2)
        assert backend._workbook is not None
        populate_table(backend._workbook, "T2", [{"Name": "second", "Value": 2}])
        backend.save(path2)
        backend.close()

        # Verify cycle 2
        wb2 = openpyxl.load_workbook(path2)
        try:
            assert wb2["Sheet1"]["A2"].value == "second"
            assert wb2["Sheet1"]["B2"].value == 2
        finally:
            wb2.close()


# ═══════════════════════════════════════════════════════════════════════════════
# 10.5 Win32Com backend (CI-safe checks)
# ═══════════════════════════════════════════════════════════════════════════════


class TestWin32ComBackend:
    """Win32Com backend: lazy Excel launch, non-Windows guard."""

    def test_win32com_backend_does_not_launch_excel_on_init(self) -> None:
        """Creating a Win32ComBackend does not start Excel."""
        from fillxl._backends._win32com import Win32ComBackend

        backend = Win32ComBackend()
        assert backend._app is None
        assert backend._workbook is None

    def test_win32com_backend_raises_on_non_windows_platform(
        self,
        tmp_path: Path,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """On non-Windows, load() raises RuntimeError with a clear message."""
        from fillxl._backends import _win32com

        monkeypatch.setattr(_win32com, "_is_windows", False)

        backend = _win32com.Win32ComBackend()
        with pytest.raises(RuntimeError, match="win32com"):
            backend.load(tmp_path / "fake.xlsx")


# ── Helpers ──────────────────────────────────────────────────────────────────


def _make_simple_table(path: Path, table_name: str) -> None:
    """Minimal workbook with one table."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    from openpyxl.worksheet.table import Table, TableColumn

    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "old"
    ws["B2"] = 0
    table = Table(
        displayName=table_name,
        ref="A1:B2",
        tableColumns=[
            TableColumn(id=1, name="Name"),
            TableColumn(id=2, name="Value"),
        ],
    )
    ws.add_table(table)
    wb.save(path)
    wb.close()
